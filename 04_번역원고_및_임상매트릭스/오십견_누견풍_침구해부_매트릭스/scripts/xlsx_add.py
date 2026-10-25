import sys,json,numpy as np,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
D=sys.argv[1]
R=json.load(open('fs_results.json',encoding='utf-8')); codes=R['codes']; F=R['features']
M=np.loadtxt('fs_similarity.csv',delimiter=',',skiprows=1)
wb=openpyxl.load_workbook(D+'/오십견_누견풍_해부임상_매트릭스.xlsx')
hdr,fill=Font(bold=True,color='FFFFFF'),PatternFill('solid',fgColor='8E44AD')
def sheet(name,head,rows,w):
    if name in wb.sheetnames: del wb[name]
    ws=wb.create_sheet(name); ws.append(head)
    for c in ws[1]: c.font,c.fill=hdr,fill; c.alignment=Alignment(wrap_text=True)
    for r in rows: ws.append(r)
    for i,x in enumerate(w,1): ws.column_dimensions[ws.cell(1,i).column_letter].width=x
sheet('정제1_FMA특징',['코드','혈명','분절','어깨관절 지배신경','더부신경','신경 영역','FMA 구조 수','다분절 제외'],
      [[c,R['names'][c][0],','.join(F[c]['G']),', '.join(F[c]['shoulder_joint_nerves']),'예' if F[c]['XI'] else '',
        '; '.join(F[c]['N']),len(F[c]['S']),'; '.join(F[c]['multiseg'])] for c in codes],[8,8,26,30,8,80,8,20])
sheet('정제2_유사도',['']+codes,[[c]+[round(float(x),3) for x in M[i]] for i,c in enumerate(codes)],[8]+[7]*len(codes))
rows=[]
for m,d in R['kinesiology'].items():
    for c,v in d['points'].items():
        rows.append([m,c,R['names'][c][0],', '.join(v['agonist']),', '.join(v['restraint']),'예' if c in d['report_points'] else '',d['capsule']])
sheet('정제3_운동기능해부',['운동','코드','혈명','주동근(A)','제한 근육(R)','원 보고서 배정','관절낭 제한(B1)'],rows,[16,8,8,40,40,10,40])
rows=[[v['label'],v['fma_id'],v['volume_cm3'],'×'.join(str(round(x)) for x in v['extent_mm']),v['watertight']] for v in R['mesh']['volumes']]
rows+=[[f"{v['point']} — {v['label']}",v['fma_id'],None,f"첫째갈비뼈 {v['Right first rib']} mm / 둘째갈비뼈 {v['Right second rib']} mm",None] for v in R['mesh']['lung_distance']]
rows+=[['주의','','','폐 메쉬(7309·7333)는 기관지 나무뿐 — 흉막·폐 표면 거리 측정 불가, 갈비뼈는 대리 지표','']]
sheet('정제4_3D측정',['구조/혈','FMA','부피(cm³)','크기·거리','닫힘'],rows,[44,8,10,60,8])
rows=[['PMID 23620934 = papers.id 1836','中国针灸 2013 — 경혈 마사지 > 전침·TENS','기전 근거 아님'],
      ['papers.id 397','J Pain Res 2023 — 침도 후향 연구','기전 근거 아님'],
      ['PMID 17333364','J Assist Reprod Genet 2007 — 자궁내막증 마이크로어레이','무관(잘못된 PMID)'],
      ['PMID 25069201','针刺研究 2014 — 뇌졸중 어깨-손 증후군 外關 온침','天宗 GTO 기전과 무관'],
      ['PMID 32952592 = papers.id 1819','eCAM 2020 — 건강 여성 근지구력 예비 연구','雲門 기전과 무관'],
      ['papers.id 1827','中国针灸 2026 — 天府 어깨 주치 교재 고찰','雲門 기전 근거 아님'],
      ['9절 SMD −0.85','출처 없음','Cochrane 2005(PMID 15846753): 근거 부족·단기 이득 가능성'],
      ['《逢源》 肩節如膠','证治通鉴_21 L1303(《针灸逢源·症治要穴歌》 표제)','출전 표기 확정·원천 미확정(허임 《鍼灸經驗方》 대조 필요)'],
      ['條口 계보','针灸腧穴通考_08 L1131~1153','일치(왕문면→王文綿)'],
      ['《席弘賦》·《世醫得效方》·《類經圖翼》','证治通鉴_21 L1043·L1213·L1221·L1263','일치'],
      ['《玉龍歌》 肩端紅腫','证治通鉴_21','92% 일치(중간 생략)']]
sheet('정제5_인용검증',['인용','실제','판정'],rows,[34,60,50])
sheet('정제6_이력',['일자','작업','도구'],[['2026-10-25','최초 작성','Antigravity'],['2026-10-25','FMA 분석·이미지 추가, 인용·운동학·Layer A·ARI 감사','Claude Code']],[12,80,16])
wb.save(D+'/오십견_누견풍_해부임상_매트릭스.xlsx'); print(wb.sheetnames)
