# -*- coding: utf-8 -*-
"""[경혈해부] 해부-임상 매트릭스 xlsx + 보고서용 표(md 조각) 생성.

사용: python build_matrix.py <출력 xlsx 경로>
"""
import csv
import json
import re
import sys
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.cluster.hierarchy import fcluster

HERE = Path(__file__).resolve().parent
OUT = Path(sys.argv[1])
IND = json.loads((HERE / "indications.json").read_text(encoding="utf-8"))
CL = json.loads((HERE / "clusters.json").read_text(encoding="utf-8"))
codes, Z, FEAT, INFO = CL["codes"], np.array(CL["linkage"]), CL["features"], CL["info"]
M = np.loadtxt(HERE / "similarity.csv", delimiter=",", skiprows=1)
RANK = list(csv.DictReader(open(HERE / "ranking.csv", encoding="utf-8-sig")))
LA = {r["code"]: r for r in csv.DictReader(open(HERE / "layerA_counts.csv", encoding="utf-8-sig"))}
R3D = json.loads((HERE / "render3d_ko.json").read_text(encoding="utf-8"))

L5, L12, L16 = (fcluster(Z, k, "maxclust") for k in (8, 12, 16))


def _name(s):
    """군집 구성원으로 층위 이름을 붙인다(보고서 6절과 같은 이름)."""
    rules = [("LU3", "A1a 수태음 본경 사슬(天府~太淵)"), ("HT7", "A1b 손목 자쪽(陰郄·神門)"), ("LU11", "A1c 손가락(少商·三間·前谷·商陽)"),
             ("SP18", "A2a 앞가슴벽 2~5늑간"), ("KI27", "A2b 빗장 아래(俞府)"), ("LU1", "A2c 앞가슴벽 바깥(中府)"), ("GB23", "A2d 옆가슴(輒筋)"),
             ("BL13", "B1 등 흉추 옆·어깨 위"), ("CV22", "B2a 흉골위패임·흉골자루"), ("ST12", "B2b 목 앞·빗장위"),
             ("GB41", "C1 발등(足臨泣)"), ("KI4", "C2 발목 안쪽(大鐘)"), ("CV15", "D 명치(鳩尾)"), ("CV18", "D 흉골 정중(玉堂)"),
             ("LR14", "D 갈비활(期門)"), ("GV14", "D C7 극돌기(大椎)")]
    hits = [n for c, n in rules if c in s]
    return " / ".join(hits) if len(hits) <= 2 else hits[0][:2] + " " + {"A1": "상지(수태음 본경·손목·손가락)", "A2": "앞가슴벽(늑간·가슴근신경)"}.get(hits[0][:2], "")


def layer_names():
    out = []
    for L in (L5, L12, L16):
        g = {}
        for c, a in zip(codes, L):
            g.setdefault(int(a), set()).add(c)
        names = {}
        for k, s in g.items():
            if "LU9" in s and "SP18" in s:
                names[k] = "A 팔신경얼기 영역(상지·앞가슴벽)"
            elif "LU9" in s and "LU11" in s:
                names[k] = "A1 상지(수태음 본경·손목·손가락)"
            elif "SP18" in s and "KI27" in s:
                names[k] = "A2 앞가슴벽(늑간·가슴근신경)"
            elif "CV22" in s and "ST12" in s:
                names[k] = "B2 목 앞·빗장위"
            elif "GB41" in s and "KI4" in s:
                names[k] = "C 하지 원위"
            else:
                names[k] = _name(s)
        out.append(names)
    return out


N5, N12, N16 = layer_names()
TRAD = {}
for line in (HERE / "make_figures.py").read_text(encoding="utf-8").splitlines():
    for a, b in re.findall(r'"([^"]+)": "([^"]+)"', line):
        if len(a) <= 3 and len(b) <= 3:
            TRAD[a] = b
GLAB = {"A_胀满膨膨": "A 肺脹滿·膨膨", "B_喘咳": "B 喘咳", "C_缺盆中痛": "C 缺盆中痛", "D_交两手而瞀_臂厥": "D 交兩手而瞀·臂厥",
        "E_心痛_心膨膨": "E 心痛(《陰陽十一脈》 원형)"}
TIERN = {"I": "1층", "II": "2층", "III": "3층"}
VISC = set(f"T{i}" for i in range(1, 6))
ARM = {"C8", "T1"}
NECK = {"C3", "C4"}


def evidence(code):
    d = IND[code]
    rows = []
    for it in d["items"]:
        gs = [g for g in it["groups"] if g != "B_喘咳"] + (["B_喘咳"] if "B_喘咳" in it["groups"] else [])
        gc = it.get("groups_combo", [])
        if not gs and not gc:
            continue
        rows.append({"line": it["line"], "era": it["era"], "source": it["source"], "groups": gs, "groups_combo": gc,
                     "flag": it.get("flag", ""), "text": it["text"], "note": it["note"]})
    return d["file"], rows


def short(t, n=70):
    return t if len(t) <= n else t[:n] + "…"


def main():
    wb = Workbook()
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F6FB2")

    def sheet(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = hdr, fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "B2"
        return ws

    wb.remove(wb.active)
    sheet("1_是動病_층차", ["층", "문헌", "원문", "근거 위치", "비고"], [
        ["원형", "《陰陽十一脈灸經》", "是动则病心滂滂如痛，缺盆痛，甚则交两手而战，此为臂厥", "针灸腧穴通考_03_上_手太阴肺经穴.md L19", "황룡상: 관상동맥 심교통 증상군"],
        ["개작", "《靈樞·經脈》", "是动则病肺胀满膨膨而喘咳，缺盆中痛，甚则交两手而瞀，此为臂厥", "같은 파일 L31, 표 2-1(L39)", "心→肺 개작으로 증상 3개가 독립 증상처럼 분리"],
        ["경혈 주치", "《黃帝明堂經》", "经渠：心痛欲呕，胸中彭彭然，甚则交两手務。咳上气喘。太渊：心痛胀满彭彭，臂厥。胸痛，缺盆中相引痛。喘不得息。尺泽：心膨膨痛。",
         "같은 파일 표 2-1(L39)", "한대 경혈 주치에는 '肺胀满膨膨' 문구가 아직 없음"],
        ["처방", "《黃帝明堂經》(《甲乙》 권8·9)", "胸中膨膨然，甚则交两手而瞀，暴瘅内逆，刺经渠及天府，此谓之大输", "针灸腧穴通考_02_上_序·凡例·总论-2.md L1027; 中国针灸证治通鉴_03 L264", "脈口(經渠)+標脈(天府) 처방 — 배오"],
    ], [10, 22, 60, 40, 40])

    rows2 = []
    for c in codes + list(CL["excluded"]):
        if c in IND:
            d = IND[c]
            rows2.append([c, TRAD.get(d["name"], d["name"]), d["pinyin"], TIERN.get(INFO.get(c, {}).get("tier", ""), "제외"),
                          " ".join(GLAB[g] for g in d["groups_single_han"]), " ".join(GLAB[g] for g in d["groups_single_all"]),
                          " ".join(GLAB[g] for g in d.get("groups_combo", [])), CL["excluded"].get(c, INFO.get(c, {}).get("note", "")),
                          d["file"]])
    sheet("2_후보경혈_주치층", ["코드", "혈명", "병음", "주치 층", "한대 단독 증상군", "전 시기 단독 증상군", "배오 구절 증상군", "판정 메모", "《通考》 파일"],
          rows2, [7, 8, 12, 7, 34, 40, 30, 50, 34])

    rows3 = []
    for c in codes + list(CL["excluded"]):
        if c not in IND:
            continue
        f, ev = evidence(c)
        for e in ev:
            rows3.append([c, e["line"], e["era"], e["source"], " ".join(GLAB[g] for g in e["groups"]),
                          " ".join(GLAB[g] for g in e["groups_combo"]), e["flag"], e["text"], e["note"][:300], f])
    sheet("3_근거조문", ["코드", "행", "시기", "출전", "단독 증상군", "배오 구절 증상군", "Rank 4 표지", "조문", "《通考》 按語", "파일"],
          rows3, [7, 6, 11, 26, 30, 22, 12, 70, 60, 34])

    rows4 = []
    for c in codes:
        f = FEAT[c]
        rows4.append([c, TRAD.get(INFO[c]["name"], INFO[c]["name"]), f["region"], ", ".join(f["G"]), ", ".join(f["cranial"]),
                      "; ".join(f["N"]), len(f["S"]), "; ".join(f.get("multiseg_excluded", [])),
                      "; ".join(f"{g}:{'/'.join(s)}" for g, s in f["G_src"].items())])
    sheet("4_해부특징_FMA", ["코드", "혈명", "부위", "척수 분절", "뇌신경", "신경 영역(FMA 라벨)", "FMA 구조 수", "다분절 근육 제외", "분절 출처"],
          rows4, [7, 8, 18, 30, 10, 80, 8, 30, 60])

    ws = wb.create_sheet("5_유사도행렬")
    ws.append([""] + codes)
    for i, c in enumerate(codes):
        ws.append([c] + [round(float(x), 3) for x in M[i]])

    rows6 = []
    for i, c in enumerate(codes):
        rows6.append([c, TRAD.get(INFO[c]["name"], INFO[c]["name"]), TIERN[INFO[c]["tier"]], N5[int(L5[i])], N12[int(L12[i])], N16[int(L16[i])]])
    sheet("6_층위", ["코드", "혈명", "주치 층", "대분류(8군)", "중분류(12군)", "소분류(16군)"], rows6, [7, 8, 7, 30, 30, 34])

    rows7 = [[r["rank_LU9"], r["code"], TRAD.get(r["name"], r["name"]), TIERN[r["tier"]], r["sim_LU9"], r["sim_LU9_eq"], r["centrality_tierI"],
              r["region"], r["segments"], r["cranial"], r["nearest3"]] for r in RANK]
    sheet("7_太淵기준_순위", ["순위", "코드", "혈명", "주치 층", "太淵 유사도(가중)", "太淵 유사도(등가중)", "1층 중심성", "부위", "분절", "뇌신경", "가장 가까운 3혈"],
          rows7, [6, 7, 8, 7, 10, 10, 10, 18, 30, 10, 40])

    rows8, seg_md = [], []
    for k in sorted(set(int(x) for x in L16)):
        mem = [c for c, l in zip(codes, L16) if int(l) == k]
        v = sum(1 for c in mem if VISC & set(FEAT[c]["G"]))
        a = sum(1 for c in mem if ARM & set(FEAT[c]["G"]))
        nk = sum(1 for c in mem if NECK & set(FEAT[c]["G"]))
        vg = sum(1 for c in mem if "CN X" in FEAT[c]["cranial"])
        idx = [codes.index(c) for c in mem]
        within = float(np.mean([M[x, y] for x in idx for y in idx if x < y])) if len(idx) > 1 else float("nan")
        row = [N16[k], len(mem), ", ".join(mem), f"{v}/{len(mem)}", f"{a}/{len(mem)}", f"{nk}/{len(mem)}", f"{vg}/{len(mem)}",
               "" if np.isnan(within) else round(within, 3)]
        rows8.append(row)
        seg_md.append(row)
    sheet("8_분절정합_B-Check②", ["소분류", "혈 수", "구성혈", "T1–T5 내장분절 포함", "C8–T1(위팔 안쪽 연관통)", "C3–C4(빗장위)", "미주신경(CN X)", "군집 내 평균 유사도"],
          rows8, [24, 6, 50, 14, 16, 12, 12, 12])

    rows9 = [[c, LA[c]["n_active"], LA[c]["n_cardiac"], LA[c]["n_resp"], LA[c]["cardiac_ids"], LA[c]["resp_ids"]] for c in codes]
    sheet("9_LayerA_DB연결", ["코드", "연결 논문(active)", "심장 관련", "호흡 관련", "심장 papers.id", "호흡 papers.id"], rows9, [7, 10, 10, 10, 40, 40])

    rows10 = [["가중치(잠정)", json.dumps(CL["weights"])], ["등가중치 민감도", json.dumps(CL["weights_eq"])],
              ["군집법", "UPGMA(average linkage), 거리 = 1 − 유사도"], ["코페네틱 상관", CL["cophenetic"]],
              ["실루엣 최적 k", CL["k"]], ["가중치 vs 등가중치 군집 일치(ARI)", CL["ari_weight_sensitivity"]],
              ["대응 수준", "FMA exact·exact+·exact(set)·parent(branch_of)만 사용, fuzzy·set(검토)·none 제외"],
              ["다분절 근육", "근육 1개가 5개 분절 이상을 끌어오면 신경·분절 특징에서 제외(래퍼 5-1절 acupoints_by_segment 규칙)"],
              ["분절 보완", "FMA 분절 미기재 신경은 B1-01(Gray's 42e)·B1-03(Moore 8e) 규범 분절로 보완"],
              ["제외", json.dumps(CL["excluded"], ensure_ascii=False)],
              ["3D", json.dumps({k: [s['label'] for s in v['structures']] for k, v in R3D.items()}, ensure_ascii=False)]]
    sheet("10_방법_감사", ["항목", "값"], rows10, [30, 120])
    wb.save(OUT)
    (HERE / "seg_table.json").write_text(json.dumps(seg_md, ensure_ascii=False), encoding="utf-8")
    print("[xlsx]", OUT)
    for r in seg_md:
        print(r)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
