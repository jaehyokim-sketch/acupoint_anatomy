# -*- coding: utf-8 -*-
"""
fma_build.py — FMA OWL 덤프 다운로드·로컬 파싱·색인·경혈 대응표·Neo4j CSV 생성 (Claude Code 전용)

사용:
  python .claude/skills/acupoint-anatomy/scripts/fma/fma_build.py status
  python .claude/skills/acupoint-anatomy/scripts/fma/fma_build.py all          # download → load → index → map → neo4j-csv
  python .claude/skills/acupoint-anatomy/scripts/fma/fma_build.py map          # 교재 xlsx가 바뀌었을 때 대응표만 재생성

단계:
  download   fma.owl(FMA 5.x, ~200 MB)을 FMA_DATA_DIR에 받는다 (이미 있으면 건너뜀, --force로 재다운로드)
  load       owlready2로 OWL/RDF를 파싱해 영속 SQLite 저장소(fma_quadstore.sqlite3)에 적재 (~30초)
  index      저장소에서 entity / term / edge 압축 색인(fma_index.sqlite3)을 SQL로 추출
  map        공통교재·김재효 교재의 '해부구조_개별항목_상세' 영문 용어를 FMA ID에 대응 → 02_Master_Datasets/FMA/
  neo4j-csv  neo4j-admin import / LOAD CSV용 nodes·relationships CSV 생성
"""
import argparse
import csv
import io
import json
import re
import shutil
import sqlite3
import sys
import time
import urllib.request
from datetime import datetime
from difflib import SequenceMatcher, get_close_matches

from fma_common import (FMA_NS, FMA_URL, INDEX_PATH, MAPPING_DIR, MAPPING_JSON, MAPPING_XLSX, NEO4J_DIR, OWL_PATH,
                        QUADSTORE_PATH, ROOT, TEXTBOOKS, ORDINALS, NORM_VERSION, data_dir, norm,
                        open_index)

RDF_TYPE = "http://www.w3.org/1999/02/22-rdf-syntax-ns#type"
OWL_CLASS = "http://www.w3.org/2002/07/owl#Class"
SUBCLASS = "http://www.w3.org/2000/01/rdf-schema#subClassOf"
ON_PROP = "http://www.w3.org/2002/07/owl#onProperty"
SOME = "http://www.w3.org/2002/07/owl#someValuesFrom"
LABEL = "http://www.w3.org/2000/01/rdf-schema#label"
ANN_SOURCE = "http://www.w3.org/2002/07/owl#annotatedSource"
TERM_PROPS = {"label": LABEL, "preferred": FMA_NS + "preferred_name", "synonym": FMA_NS + "synonym",
              "latin": FMA_NS + "non-English_equivalent"}


# ---------------------------------------------------------------- download
def step_download(force=False):
    if OWL_PATH.exists() and not force:
        print(f"[download] 기존 파일 사용: {OWL_PATH} ({OWL_PATH.stat().st_size // 2**20} MB)")
        return
    part = OWL_PATH.with_suffix(".owl.part")
    print(f"[download] {FMA_URL} → {OWL_PATH}")
    t = time.time()
    with urllib.request.urlopen(FMA_URL, timeout=120) as r, open(part, "wb") as f:
        shutil.copyfileobj(r, f, length=1 << 20)
    head = open(part, "rb").read(4096).decode("utf-8", "ignore")
    if "fma.owl" not in head or "owl:Ontology" not in head:
        part.unlink()
        sys.exit("[download] 받은 파일이 FMA OWL이 아닙니다 (리다이렉트 오류 가능). URL을 확인하세요.")
    part.replace(OWL_PATH)
    print(f"[download] 완료 {OWL_PATH.stat().st_size // 2**20} MB, {time.time() - t:.0f}s")


# ---------------------------------------------------------------- load (owlready2)
def step_load(force=False):
    import owlready2
    if QUADSTORE_PATH.exists():
        if not force and QUADSTORE_PATH.stat().st_mtime >= OWL_PATH.stat().st_mtime:
            print(f"[load] 기존 저장소 사용: {QUADSTORE_PATH}")
            return
        QUADSTORE_PATH.unlink()
    t = time.time()
    owlready2.default_world.set_backend(filename=str(QUADSTORE_PATH), exclusive=False)
    onto = owlready2.get_ontology("file://" + str(OWL_PATH).replace("\\", "/")).load()
    owlready2.default_world.save()
    print(f"[load] owlready2 적재 완료 {time.time() - t:.0f}s, 저장소 {QUADSTORE_PATH.stat().st_size // 2**20} MB, "
          f"version={getattr(onto.metadata, 'versionInfo', ['?'])}")


# ---------------------------------------------------------------- index
def step_index():
    t = time.time()
    q = sqlite3.connect(str(QUADSTORE_PATH))
    sid = {iri: s for s, iri in q.execute("SELECT storid, iri FROM resources")}
    iri_of = {s: iri for iri, s in sid.items()}

    def fid(storid):
        iri = iri_of.get(storid, "")
        if iri.startswith(FMA_NS + "fma") and iri[len(FMA_NS) + 3:].isdigit():
            return int(iri[len(FMA_NS) + 3:])
        return None

    classes = {s for (s,) in q.execute("SELECT s FROM objs WHERE p=? AND o=?", (sid[RDF_TYPE], sid[OWL_CLASS]))}
    ent, seen = {}, {}
    # 원본에 'fma0323167'처럼 0이 덧붙은 IRI가 있어 숫자 ID가 겹친다 → 표준 IRI(0 없음)를 대표로 삼고 나머지는 같은 레코드로 합친다
    for s in sorted(classes, key=lambda s: iri_of.get(s, "")[len(FMA_NS) + 3:].startswith("0")):
        f = fid(s)
        if f is None:
            continue
        if f in seen:
            ent[s] = ent[seen[f]]
        else:
            seen[f] = s
            ent[s] = {"fma_id": f, "label": None, "preferred": None, "definition": None, "ta_id": None}

    terms = []
    for kind, prop in TERM_PROPS.items():
        if prop not in sid:
            continue
        for s, o in q.execute("SELECT s, o FROM datas WHERE p=?", (sid[prop],)):
            if s in ent and isinstance(o, str):
                terms.append((norm(o), ent[s]["fma_id"], kind, o))
                if kind in ("label", "preferred") and not ent[s][kind]:
                    ent[s][kind] = o
    if FMA_NS + "definition" in sid:
        for s, o in q.execute("SELECT s, o FROM datas WHERE p=?", (sid[FMA_NS + "definition"],)):
            if s in ent:
                ent[s]["definition"] = o
    # TA_ID는 라틴명 주석 공리(owl:Axiom)에 붙어 있다
    if FMA_NS + "TA_ID" in sid:
        for s, ta in q.execute("""SELECT a.o, d.o FROM objs a JOIN datas d ON d.s = a.s
                                  WHERE a.p=? AND d.p=?""", (sid[ANN_SOURCE], sid[FMA_NS + "TA_ID"])):
            if s in ent and not ent[s]["ta_id"]:
                ent[s]["ta_id"] = ta

    edges = set()
    for s, o in q.execute("SELECT s, o FROM objs WHERE p=? AND o>0", (sid[SUBCLASS],)):
        if s in ent and o in ent:
            edges.add((ent[s]["fma_id"], "is_a", ent[o]["fma_id"]))
    for s, prop, o in q.execute("""SELECT sc.s, op.o, sv.o FROM objs sc
                                   JOIN objs op ON op.s = sc.o AND op.p = ?
                                   JOIN objs sv ON sv.s = sc.o AND sv.p = ?
                                   WHERE sc.p = ? AND sc.o < 0""", (sid[ON_PROP], sid[SOME], sid[SUBCLASS])):
        if s in ent and o in ent:
            rel = iri_of.get(prop, "").replace(FMA_NS, "")
            edges.add((ent[s]["fma_id"], rel, ent[o]["fma_id"]))
    q.close()

    tmp = INDEX_PATH.with_suffix(".tmp")
    if tmp.exists():
        tmp.unlink()
    con = sqlite3.connect(str(tmp))
    con.executescript("""
        CREATE TABLE entity (fma_id INTEGER PRIMARY KEY, label TEXT, preferred TEXT, ta_id TEXT, definition TEXT);
        CREATE TABLE term (term_norm TEXT, fma_id INTEGER, kind TEXT, term TEXT);
        CREATE TABLE edge (src INTEGER, rel TEXT, dst INTEGER);
        CREATE TABLE meta (key TEXT PRIMARY KEY, value TEXT);
    """)
    con.executemany("INSERT INTO entity VALUES (?,?,?,?,?)",
                    [(e["fma_id"], e["label"] or e["preferred"], e["preferred"], e["ta_id"], e["definition"])
                     for e in {e["fma_id"]: e for e in ent.values()}.values()])
    con.executemany("INSERT INTO term VALUES (?,?,?,?)", terms)
    con.executemany("INSERT INTO edge VALUES (?,?,?)", sorted(edges))
    con.executescript("""
        CREATE INDEX ix_term ON term(term_norm);
        CREATE INDEX ix_term_id ON term(fma_id);
        CREATE INDEX ix_edge_src ON edge(src, rel);
        CREATE INDEX ix_edge_dst ON edge(dst, rel);
    """)
    head = open(OWL_PATH, "rb").read(3000).decode("utf-8", "ignore")
    ver = re.search(r"<owl:versionInfo[^>]*>([^<]+)<", head)
    con.executemany("INSERT INTO meta VALUES (?,?)", [("fma_version", ver.group(1) if ver else "?"),
                                                     ("built_at", datetime.now().isoformat(timespec="seconds")),
                                                     ("norm_version", NORM_VERSION)])
    con.commit()
    con.close()
    tmp.replace(INDEX_PATH)
    print(f"[index] entity {len({e['fma_id'] for e in ent.values()}):,} / term {len(terms):,} / edge {len(edges):,}  ({time.time() - t:.0f}s) → {INDEX_PATH}")


# ---------------------------------------------------------------- map (교재 영문 용어 → FMA)
HAND_MERIDIANS = ("LU", "LI", "HT", "SI", "PC", "TE")
# 교재 표기 → FMA 표기 (정규화 후 문자열 기준)
ALIASES = {
    r"^(external (abdominis |abdominal )?oblique( abdominis)?|external abdominis oblique|obliquus abdominis externus)$": "external oblique",
    r"^(internal (abdominis |abdominal )?oblique( abdominis)?|internal abdominis oblique|obliquus abdominis internus)$": "internal oblique",
    r"^abdominal transverse$": "transversus abdominis",
    r"^intercostal artery$": "posterior intercostal artery",
    r"^intercostal vein$": "posterior intercostal vein",
    r"^rectus posterior major capitis$": "rectus capitis posterior major",
    r"^rectus posterior minor capitis$": "rectus capitis posterior minor",
    r"^levator scapular$": "levator scapulae",
    r"^iliocostal$": "iliocostalis",
    r"^frontal belly( of occipitofrontalis)?$": "frontalis",
    r"^(semimembranous)$": "semimembranosus",
    r"^(semitendinous)$": "semitendinosus",
    r"^transversus thoracic$": "transversus thoracis",
    r"\bcluneal\b": "clunial",
    r"\binterossei dorsales\b": "dorsal interosseous",
    r"\binterossei palmares\b": "palmar interosseous",
    r"\bdorsal ve\w*us (plexus|network) of hand\b": "dorsal venous network of hand",
    r"^anterior interosseous nerve of forearm$": "anterior interosseous nerve",
    r"^ulnar artery vein$": "ulnar vein",
    # 척수신경 가지: 교재 "posterior branch of the 5th thoracic nerve" ↔ FMA "Posterior ramus of fifth thoracic nerve"
    r"\b(posterior|anterior|dorsal|ventral) branch of ((?:\w+ )?(?:cervical|thoracic|lumbar|sacral|spinal) nerve)\b": r"\1 ramus of \2",
    r"\b(medial|lateral) branch of posterior branch of\b": r"\1 branch of posterior ramus of",
}
# 근사 대응에서 서로 바뀌면 다른 구조가 되는 대립어
OPPOSITES = [{"palmar", "plantar"}, {"anterior", "posterior"}, {"medial", "lateral"}, {"superior", "inferior"},
             {"dorsal", "ventral"}, {"superficial", "deep"}, {"left", "right"}, {"greater", "lesser"},
             {"internal", "external"}, {"hand", "foot"}, {"antebrachial", "femoral"}, {"major", "minor"},
             {"proximal", "distal"}, {"ascending", "descending"}]
SEGMENT_NAMES = {"C": "cervical", "T": "thoracic", "L": "lumbar", "S": "sacral"}
HEAD_NOUNS = ("nerve", "artery", "vein", "muscle", "ligament", "fascia", "plexus", "branch", "tendon", "bone",
              "retinaculum", "aponeurosis", "sheath", "membrane", "ramus")


def _spelling_variant(w, x):
    """오탈자 수준 차이인가 (tempotalis~temporalis, ulna~ulnar). 접두어로 다른 단어(occipital~suboccipital)나 첫 글자가 다르면 거부."""
    if w[0] != x[0]:
        return False
    if (w in x or x in w) and abs(len(w) - len(x)) >= 3:
        return False
    return SequenceMatcher(None, w, x).ratio() >= 0.8


def _head(n):
    for w in reversed(n.split()):
        if w in HEAD_NOUNS:
            return w
    return None


class Matcher:
    def __init__(self, con):
        self.con = con
        self.by_term, self.by_squash = {}, {}
        for tn, f, label in con.execute("SELECT t.term_norm, t.fma_id, e.label FROM term t JOIN entity e USING(fma_id)"):
            self.by_term.setdefault(tn, {})[f] = label
            self.by_squash.setdefault(tn.replace(" ", ""), {})[f] = label
        self.keys = list(self.by_term)

    def _pick(self, cands: dict, region=None):
        # 좌우 미지정 실체 우선 → (부위 단서 일치) → 짧은 라벨 우선
        def key(item):
            f, label = item
            low = (label or "").lower()
            off_region = bool(region) and not low.endswith(region)
            return (low.startswith(("left ", "right ")) or " left " in low or " right " in low, off_region, len(low))
        f, label = sorted(cands.items(), key=key)[0]
        return f, label

    @staticmethod
    def _alias(n):
        for pat, rep in ALIASES.items():
            if re.search(pat, n):
                if pat.startswith("^"):
                    return rep
                n = re.sub(pat, rep, n)
        return n

    def _exact(self, text, region=None):
        n = self._alias(norm(text))
        plural = re.sub(r"(\w+)$", lambda m: m.group(1) + ("es" if m.group(1).endswith("s") else "s"), n)
        cands = [(n, "exact"), (n + " muscle", "exact+muscle"), (n + " nerve", "exact+nerve"),
                 (re.sub(r"\bmuscle$", "", n).strip(), "exact-muscle"), ("set of " + plural, "exact(set)")]
        if region:
            cands += [(f"{n} of {r}", "exact+region") for r in region]
        for cand, how in cands:
            pool = self.by_term.get(cand) or self.by_squash.get(cand.replace(" ", ""))
            if pool:
                f, label = self._pick(pool, region[0] if region and how != "exact+region" else None)
                # 일반 용어("dorsal digital nerve")가 특정 기원 신경의 집합("... of radial nerve")에 동의어로만 걸린 경우는 확정하지 않는다
                lab_n = norm(label or "")
                m = re.search(r"\bof (\w+ nerve)$", lab_n)
                if lab_n.startswith("set of ") and m and m.group(1) not in n:
                    how = "set(검토)" if how in ("exact", "exact(set)") else how + "(검토)"
                return f, label, how
        return None

    def _fuzzy(self, text):
        n = self._alias(norm(text))
        close = get_close_matches(n, self.keys, n=8, cutoff=0.88)
        ords = {w for w in n.split() if w in ORDINALS}
        for c in close:
            # 서수(분절 번호)·핵심 명사(nerve/artery/vein…)가 다르거나 대립어가 바뀌면 거부
            if {w for w in c.split() if w in ORDINALS} != ords or _head(c) != _head(n):
                continue
            a, b = set(n.split()), set(c.split())
            if any(len(pair & a) == 1 and len(pair & b) == 1 and pair & a != pair & b for pair in OPPOSITES):
                continue
            # 서로 다른 단어는 철자 변이 수준(유사도 ≥ 0.8)이어야 한다
            if not all(any(_spelling_variant(w, x) for x in b - a) for w in a - b):
                continue
            if len(a - b) != len(b - a):
                continue
            return self._pick(self.by_term[c])
        return None

    @staticmethod
    def _variants(text):
        t = text.replace("ﾠ", " ").strip()
        parent = None
        m = re.search(r"\(from ([^)]+)\)|\bfrom (?:the )?([a-z][a-z ]+?(?:nerve|artery|vein))\b", t, re.I)
        if m:
            parent = m.group(1) or m.group(2)
        t = re.sub(r"\([^)]*\)", " ", t)
        t = re.sub(r"\s*(=|;).*$", "", t)                                   # "= 별칭", "; CN5" 꼬리 제거
        t = re.sub(r"\bfrom (?:the )?\S.*$", "", t, flags=re.I).strip()       # "from C5-C7", "from the X" 꼬리 제거
        if "," in t:                                                          # "median n., proper palmar digital nerve"
            parts = [x.strip() for x in t.split(",") if not re.fullmatch(r"[CTLS]\d{1,2}(\s*-\s*[CTLS]?\d{1,2})?", x.strip())]
            t = parts[-1] if parts else t
        t = re.sub(r"\bartery vein\b", "artery and vein", t)
        # 척수 분절 코드: "posterior ramus of T12" → "... of 12th thoracic nerve", "L1-L2" → 분절별 전개
        seg = re.search(r"\b([CTLS])(\d{1,2})(?:\s*-\s*[CTLS]?(\d{1,2}))?\b", t)
        if seg and re.search(r"\b(ramus|branch|nerve)\b", t[:seg.start()], re.I):
            lo, hi = int(seg.group(2)), int(seg.group(3) or seg.group(2))
            name = SEGMENT_NAMES[seg.group(1)]
            pre = t[:seg.start()].rstrip()
            tail = "" if pre.lower().endswith("nerve") else " nerve"
            return [f"{pre} {k}th {name}{tail}".replace(" of nerve", "") for k in range(lo, hi + 1) if 1 <= k <= 12], parent
        vs = [t]
        if re.search(r"\barter(y|ies) and veins?\b", t, re.I):
            vs = [re.sub(r"\barter(y|ies) and veins?\b", "artery", t, flags=re.I),
                  re.sub(r"\barter(y|ies) and veins?\b", "vein", t, flags=re.I)]
        # 분절 범위 전개: "10th-11th intercostal nerve", "4th ~ 5th lumbar nerve"
        out = []
        for v in vs:
            r = re.search(r"\b(\d{1,2})(?:st|nd|rd|th)?\s*(?:-|~|to)\s*(\d{1,2})(?:st|nd|rd|th)\b", v)
            if r and 1 <= int(r.group(1)) < int(r.group(2)) <= 12:
                for k in range(int(r.group(1)), int(r.group(2)) + 1):
                    out.append(v[:r.start()] + f"{k}th " + v[r.end():].lstrip())
            else:
                out.append(v)
        return out, parent

    def match(self, text, code=None):
        out = []
        if not text:
            return out
        region = None
        if code:
            region = ("hand", "wrist") if code[:2] in HAND_MERIDIANS else ("foot", "ankle")
        variants, parent_hint = self._variants(text)
        for v in variants:
            v = re.sub(r"\b(\d{1,2})th\b", lambda m: {1: "1st", 2: "2nd", 3: "3rd"}.get(int(m.group(1)), m.group(0)), v)
            # 손등 'dorsal digital nerve (from radial/ulnar nerve)' → FMA 'Set of dorsal digital branches of X' (일반 대조보다 먼저)
            if parent_hint and re.fullmatch(r"dorsal digital nerve", norm(v)):
                hit = self._exact(f"set of dorsal digital branches of {parent_hint}")
                if hit and norm(parent_hint) in norm(hit[1]):
                    out.append({"variant": v, "fma_id": hit[0], "fma_label": hit[1], "match": "exact(set)"})
                    continue
            hit = self._exact(v, region)
            if hit:
                out.append({"variant": v, "fma_id": hit[0], "fma_label": hit[1], "match": hit[2]})
                continue
            hit = self._fuzzy(v)
            if hit:
                out.append({"variant": v, "fma_id": hit[0], "fma_label": hit[1], "match": "fuzzy(검토)"})
                continue
            # "... branch of X" → 상위 X (branch_of 근사)
            mb = re.search(r"\b(?:branch(?:es)?|brs?\.?|ramus|rami) of (?:the )?(.+)$", v, re.I)
            target = mb.group(1) if mb else parent_hint
            hit = self._exact(target, region) if target else None
            if hit:
                out.append({"variant": v, "fma_id": hit[0], "fma_label": hit[1], "match": "parent(branch_of)"})
            else:
                out.append({"variant": v, "fma_id": None, "fma_label": None, "match": "none"})
        return out


def step_map():
    import openpyxl
    con = open_index()
    fma_ver = con.execute("SELECT value FROM meta WHERE key='fma_version'").fetchone()[0]
    matcher = Matcher(con)
    rows, cache = [], {}
    for src, (fname, sheet, eng_col) in TEXTBOOKS.items():
        path = ROOT / "02_Master_Datasets" / fname
        ws = openpyxl.load_workbook(path, read_only=True)[sheet]
        it = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(it)]
        col = {h: i for i, h in enumerate(header)}
        for r in it:
            code = r[col["경혈코드"]]
            if not code:
                continue
            eng = (r[col[eng_col]] or "").strip()
            ck = (eng, str(code)[:2] in HAND_MERIDIANS)
            if ck not in cache:
                cache[ck] = matcher.match(eng, str(code))
            for hit in cache[ck] or [{"variant": "", "fma_id": None, "fma_label": None, "match": "none"}]:
                rows.append({"source": src, "code": code, "name_kr": r[col["표준경혈명"]], "category": r[col["해부구분"]],
                             "term_kr": r[col["한글용어"]], "term_en": eng, **hit})
    con.close()

    MAPPING_DIR.mkdir(parents=True, exist_ok=True)
    stats = {}
    for row in rows:
        k = row["match"].split("(")[0]
        stats.setdefault(row["source"], {}).setdefault(k, 0)
        stats[row["source"]][k] += 1
    payload = {"_metadata": {"generated_at": datetime.now().isoformat(timespec="seconds"), "fma_version": fma_ver,
                             "generator": ".claude/skills/acupoint-anatomy/scripts/fma/fma_build.py map",
                             "match_levels": {"exact": "라벨·선호명·동의어·라틴명 정확 일치(정규화 후)",
                                              "exact+muscle/nerve/-muscle": "접미어 보완 후 정확 일치",
                                              "exact(set)": "FMA 'Set of …' 집합 실체와 일치",
                                              "exact+region": "경락 부위 단서(수경→hand/wrist, 족경→foot/ankle)로 'X of hand/foot' 선택",
                                              "parent(branch_of)": "가지 자체가 FMA에 없어 상위 신경·혈관에 대응 — 인용 시 '상위 구조 대응' 명시",
                                              "fuzzy(검토)": "유사도 ≥ 0.88 근사 대응 — 사람 검토 전 인용 금지",
                                              "none": "대응 실패"},
                             "stats": stats},
               "rows": rows}
    io.open(MAPPING_JSON, "w", encoding="utf-8").write(json.dumps(payload, ensure_ascii=False, indent=1))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "경혈_FMA_대응표"
    cols = ["source", "code", "name_kr", "category", "term_kr", "term_en", "variant", "fma_id", "fma_label", "match"]
    ws.append(["출처교재", "경혈코드", "경혈명", "해부구분", "한글용어", "영문용어", "대조어", "FMA_ID", "FMA_label", "대응수준"])
    for row in rows:
        ws.append([row[c] for c in cols])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws2 = wb.create_sheet("대응_통계")
    ws2.append(["출처교재", "대응수준", "행수"])
    for s, d in stats.items():
        for k, v in sorted(d.items()):
            ws2.append([s, k, v])
    ws3 = wb.create_sheet("미대응_용어")
    ws3.append(["영문용어", "대응수준", "출현수"])
    miss = {}
    for row in rows:
        if row["match"] in ("none",) or row["match"].startswith("fuzzy"):
            miss.setdefault((row["term_en"], row["match"]), 0)
            miss[(row["term_en"], row["match"])] += 1
    for (t, m), n in sorted(miss.items(), key=lambda x: -x[1]):
        ws3.append([t, m, n])
    out_xlsx = MAPPING_XLSX
    try:
        wb.save(out_xlsx)
    except PermissionError:   # Excel에서 열려 있으면 잠겨 있다
        out_xlsx = MAPPING_XLSX.with_name(f"{MAPPING_XLSX.stem}_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        wb.save(out_xlsx)
        print(f"[map] ⚠ {MAPPING_XLSX.name}이 열려 있어 {out_xlsx.name}로 저장했습니다. Excel을 닫고 map을 다시 실행하세요.")
    print(f"[map] {len(rows):,}행 → {MAPPING_JSON.relative_to(ROOT)}, {out_xlsx.name}")
    for s, d in stats.items():
        tot = sum(d.values())
        ok = sum(v for k, v in d.items() if k.startswith("exact"))
        print(f"      {s}: {d}  정확 대응률 {ok / tot:.1%}")


# ---------------------------------------------------------------- neo4j csv
def step_neo4j_csv():
    con = sqlite3.connect(str(INDEX_PATH))
    NEO4J_DIR.mkdir(parents=True, exist_ok=True)
    with io.open(NEO4J_DIR / "fma_nodes.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["fmaId:ID(FMA)", "label", "taId", ":LABEL"])
        for fid, label, ta in con.execute("SELECT fma_id, label, ta_id FROM entity"):
            w.writerow([fid, label, ta or "", "Anatomy"])
    # FMA에 branch_of/tributary_of가 없고 is_a 클래스명('Branch of X', 'Tributary of X')으로만 줄기가 드러나는 실체에
    # 파생 관계 BRANCH_OF_INFERRED를 만든다 (fma_query.py trunk_chain과 같은 규칙). 예: Deep fibular nerve → Common fibular nerve
    label_to_id = {}
    for tn, fid in con.execute("SELECT term_norm, fma_id FROM term WHERE kind IN ('label','preferred')"):
        label_to_id.setdefault(tn, fid)
    has_trunk = {s for (s,) in con.execute("SELECT DISTINCT src FROM edge WHERE rel IN ('branch_of','tributary_of')")}
    labels = dict(con.execute("SELECT fma_id, label FROM entity"))
    inferred = set()
    for s, d in con.execute("SELECT src, dst FROM edge WHERE rel='is_a'"):
        if s in has_trunk:
            continue
        m = re.match(r"^((?:[\w-]+ )*?(?:branch|tributary|ramus)) of (.+)$", labels.get(d) or "", re.I)
        if m:
            t = label_to_id.get(norm(m.group(2)))
            if t and t != s:
                inferred.add((s, t))
    with io.open(NEO4J_DIR / "fma_rels.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow([":START_ID(FMA)", ":END_ID(FMA)", ":TYPE"])
        for s, rel, d in con.execute("SELECT src, rel, dst FROM edge"):
            w.writerow([s, d, rel.upper()])
        for s, t in sorted(inferred):
            w.writerow([s, t, "BRANCH_OF_INFERRED"])
    print(f"[neo4j-csv] 파생 줄기 관계 BRANCH_OF_INFERRED {len(inferred):,}")
    # 경혈 노드와 경혈→FMA 대응 관계 (대응표가 있을 때)
    if MAPPING_JSON.exists():
        rows = json.loads(io.open(MAPPING_JSON, encoding="utf-8").read())["rows"]
        pts = {}
        for r in rows:
            pts.setdefault(r["code"], r["name_kr"])
        with io.open(NEO4J_DIR / "acupoint_nodes.csv", "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["code:ID(Acupoint)", "nameKr", ":LABEL"])
            for c, n in sorted(pts.items()):
                w.writerow([c, n, "Acupoint"])
        with io.open(NEO4J_DIR / "acupoint_rels.csv", "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow([":START_ID(Acupoint)", ":END_ID(FMA)", "source", "category", "match", ":TYPE"])
            seen = set()
            for r in rows:
                if r["fma_id"] and not r["match"].startswith(("fuzzy", "none")) and "검토" not in r["match"]:
                    k = (r["code"], r["fma_id"], r["source"])
                    if k not in seen:
                        seen.add(k)
                        w.writerow([r["code"], r["fma_id"], r["source"], r["category"], r["match"], "HAS_STRUCTURE"])
    con.close()
    print(f"[neo4j-csv] → {NEO4J_DIR}")
    print("   권장: 실행 중인 서버에 적재 → python .claude/skills/acupoint-anatomy/scripts/fma/fma_neo4j.py load")
    print("   대용량 초기 적재(서버 정지·빈 DB): neo4j-admin database import full --nodes=fma_nodes.csv --nodes=acupoint_nodes.csv "
          "--relationships=fma_rels.csv --relationships=acupoint_rels.csv neo4j")
    print("     (이 경로는 fmaId가 문자열로 저장되므로 Cypher 템플릿 사용 전 "
          "MATCH (n:Anatomy) SET n.fmaId = toInteger(n.fmaId) 를 한 번 실행)")


def step_status():
    print(f"[status] FMA_DATA_DIR = {data_dir()}")
    for p in (OWL_PATH, QUADSTORE_PATH, INDEX_PATH, MAPPING_JSON, MAPPING_XLSX, NEO4J_DIR / "fma_nodes.csv"):
        if p.exists():
            print(f"  [OK] {p}  {p.stat().st_size / 2**20:.1f} MB  {datetime.fromtimestamp(p.stat().st_mtime):%Y-%m-%d %H:%M}")
        else:
            print(f"  [--] {p}")
    for mod in ("rdflib", "owlready2", "neo4j", "py2neo", "openpyxl"):
        try:
            __import__(mod)
            print(f"  [OK] python: {mod}")
        except Exception:
            print(f"  [--] python: {mod}  (pip install -r .claude/skills/acupoint-anatomy/scripts/fma/requirements-fma.txt)")


def main():
    ap = argparse.ArgumentParser(description="FMA 로컬 파싱·색인·경혈 대응표 빌더")
    ap.add_argument("step", choices=["status", "download", "load", "index", "map", "neo4j-csv", "all"])
    ap.add_argument("--force", action="store_true", help="download/load 단계 재실행")
    a = ap.parse_args()
    if a.step == "status":
        return step_status()
    steps = ["download", "load", "index", "map", "neo4j-csv"] if a.step == "all" else [a.step]
    for s in steps:
        {"download": lambda: step_download(a.force), "load": lambda: step_load(a.force), "index": step_index,
         "map": step_map, "neo4j-csv": step_neo4j_csv}[s]()


if __name__ == "__main__":
    main()
