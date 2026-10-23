# -*- coding: utf-8 -*-
"""
sync_master_db.py — 경혈 생의학 통합 마스터 DB 동기화 (Master DB Sync v3.0)

Claude Code(.claude 래퍼)와 Antigravity(.agents 스킬)가 공용으로 호출하는 단일 진입점.
설계 문서: 01_Framework/biomedical_master_db_sync_design.md

서브커맨드
  scan-local   로컬 논문 폴더 증분 스캔(manifest 비교) → papers/역색인 갱신
  sync-api     PubMed · OpenAlex · Crossref (· Europe PMC 폴백) 수집 → 레지스트리 중복 제거 · 병합 적재
  export-xlsx  JSON에서 xlsx 5개 시트 전면 재생성
  validate     JSON ↔ xlsx 정합성 검증 (실패 시 종료코드 1)
  all          scan-local → sync-api → export-xlsx → validate (Rank 6 표준 호출)
  migrate-v3   v2.6 스냅샷을 v3.0 스키마로 1회 변환

원칙: JSON이 단일 진실원천, xlsx는 파생물. 논문 ID는 영구. 저장은 원자적이며 직전본을 backup/에 보관.
"""

import argparse
import ast
import csv
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import time
import urllib.parse
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    import requests
except ImportError:  # pragma: no cover
    requests = None
try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None
try:
    import olefile
except ImportError:
    olefile = None
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment
except ImportError:  # pragma: no cover
    openpyxl = None

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# --------------------------------------------------------------------------------------
# 경로 · 상수
# --------------------------------------------------------------------------------------
SCRIPT_DIR = Path(__file__).resolve().parent

def find_project_root(start_path: Path) -> Path:
    curr = start_path.resolve()
    for _ in range(6):
        if (curr / "01_Framework").exists() and (curr / "02_Master_Datasets").exists():
            return curr
        if curr.parent == curr:
            break
        curr = curr.parent
    cwd = Path.cwd().resolve()
    if (cwd / "01_Framework").exists() and (cwd / "02_Master_Datasets").exists():
        return cwd
    curr = start_path.resolve()
    for _ in range(4):
        if (curr / "framework").exists() and (curr / "data").exists():
            return curr
        if curr.parent == curr:
            break
        curr = curr.parent
    return start_path.resolve()

DEFAULT_ROOT = find_project_root(SCRIPT_DIR)

JSON_NAME = "경혈_생의학_신경반사_마스터DB.json"
XLSX_NAME = "경혈_생의학_신경반사_마스터DB.xlsx"
CSV_NAME = "361_Acupoints_Refined_Master_DB_v2.csv"
DEFAULT_LOCAL_DIR = r"G:\내 드라이브\학습자료\실험 및 연구 논문"

SHEET_MASTER = "생의학_신경반사_마스터DB"
SHEET_LOCAL_INDEX = "로컬_논문_역색인"
SHEET_PAPERS = "논문_전수_색인"
SHEET_EXTERNAL = "외부_수집_논문"
SHEET_LOG = "동기화_로그"
LEGACY_SHEETS = {"로컬_논문_역색인_1673편": SHEET_LOCAL_INDEX, "소장_논문_전수_색인_1673편": SHEET_PAPERS}

SCHEMA_VERSION = 3
DB_VERSION = "3.0_unified"
LOCAL_EXTS = {".pdf", ".hwp", ".hwpx"}
HASH_BYTES = 1024 * 1024  # 첫 1MB + 파일 크기
SYNC_LOG_CAP = 500
HTTP_TIMEOUT = 15
UA = "AcupointAnatomySync/3.0 (mailto:{})".format(os.environ.get("OPENALEX_MAILTO", "researcher@local.workspace"))
NCBI_KEY = os.environ.get("NCBI_API_KEY", "")
NCBI_SLEEP = 0.11 if NCBI_KEY else 0.35
CN_JOURNAL_PREFIXES = {"10.13703": "Zhongguo Zhen Jiu", "10.13702": "Zhen Ci Yan Jiu"}
MERIDIAN_ORDER = ["LU", "LI", "ST", "SP", "HT", "SI", "BL", "KI", "PC", "TE", "SJ", "GB", "LR", "LV", "GV", "DU", "CV", "RN", "EX"]

ORIGIN_MAP = {
    "pubmed": "pubmed", "openalex": "openalex", "openalex(cnki proxy)": "openalex", "crossref": "crossref",
    "europepmc": "europepmc", "local drive": "local", "local drive (indexed)": "local", "db": "manual", "": "manual",
}


class Ctx:
    def __init__(self, root: Path, dry_run: bool, trigger: str):
        self.root = root
        self.dry_run = dry_run
        self.trigger = trigger
        self.ds = root / "02_Master_Datasets"
        self.json_path = self.ds / JSON_NAME
        self.xlsx_path = self.ds / XLSX_NAME
        self.csv_path = self.ds / CSV_NAME
        self.backup_dir = self.ds / "backup"
        self.log_dir = self.ds / "sync_logs"
        self._dict: Optional[Dict[str, Dict[str, Any]]] = None

    def acupoint_dict(self) -> Dict[str, Dict[str, Any]]:
        if self._dict is None:
            self._dict = load_acupoint_dictionary(self.csv_path)
        return self._dict


# --------------------------------------------------------------------------------------
# 공통 유틸
# --------------------------------------------------------------------------------------
def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def std_code(code: str) -> str:
    c = (code or "").strip().upper()
    m = re.match(r"^([A-Z]+)[- ]?0*(\d+)$", c)
    return f"{m.group(1)}{int(m.group(2))}" if m else c


def normalize_doi(doi: Optional[str]) -> str:
    if not doi:
        return ""
    d = str(doi).strip().lower()
    d = re.sub(r"^https?://(dx\.)?doi\.org/", "", d)
    d = re.sub(r"^doi:\s*", "", d)
    d = d.strip().rstrip(".,;")
    # DOI 형식 검증: '10.xxxx/...' 가 아니면 DOI로 취급하지 않음 (예: 출판사명이 DOI 칸에 들어간 경우)
    return d if re.match(r"^10\.\d{4,9}/\S+$", d) else ""


def normalize_title(title: Optional[str]) -> str:
    if not title:
        return ""
    t = str(title).lower()
    return re.sub(r"[^a-z0-9\u4e00-\u9fff\uac00-\ud7af]", "", t)


def title_hash(p: Dict[str, Any]) -> str:
    t = normalize_title(p.get("title"))
    if not t:
        return ""
    first = str(p.get("authors") or "").split(",")[0].strip().lower()
    return f"{t[:40]}_{first[:10]}_{p.get('year') or ''}"


def as_list(v) -> list:
    if isinstance(v, list):
        return v
    if v in (None, "", "[]"):
        return []
    if isinstance(v, str):
        try:
            r = ast.literal_eval(v)
            return r if isinstance(r, list) else [r]
        except Exception:
            return [v]
    return [v]


def as_dict(v) -> dict:
    if isinstance(v, dict):
        return v
    if isinstance(v, str) and v.strip().startswith("{"):
        try:
            r = ast.literal_eval(v)
            return r if isinstance(r, dict) else {}
        except Exception:
            return {}
    return {}


def as_int(v, default=0) -> int:
    try:
        return int(v)
    except Exception:
        return default


def file_hash(path: Path) -> str:
    h = hashlib.sha1()
    try:
        size = os.stat(lp(path)).st_size
        with open(lp(path), "rb") as f:
            h.update(f.read(HASH_BYTES))
        h.update(str(size).encode())
        return h.hexdigest()
    except Exception:
        return ""


def relpath_norm(path: Path, base: Path) -> str:
    try:
        return os.path.relpath(str(path), str(base)).replace("\\", "/")
    except ValueError:
        return str(path).replace("\\", "/")


def lp(path) -> str:
    """Windows 260자 경로 한계 회피: 긴 절대경로에 확장 경로 접두(\\\\?\\)를 붙인다."""
    s = str(path)
    if os.name == "nt" and len(s) > 240 and not s.startswith("\\\\?\\"):
        s = "\\\\?\\" + os.path.abspath(s)
    return s


def parse_filename_meta(filename: str) -> Dict[str, Any]:
    base = re.sub(r"\.(pdf|hwpx?|PDF|HWPX?)$", "", filename)
    parts = base.split("_")
    meta = {"year": None, "journal": "", "author": "", "title": base}
    if len(parts) >= 2 and re.match(r"^(19\d\d|20\d\d)$", parts[0]):
        meta["year"] = int(parts[0])
        if len(parts) >= 4:
            meta["journal"], meta["author"], meta["title"] = parts[1], parts[2], "_".join(parts[3:])
        elif len(parts) == 3:
            meta["journal"], meta["title"] = parts[1], parts[2]
        else:
            meta["title"] = "_".join(parts[1:])
    return meta


# --------------------------------------------------------------------------------------
# 경혈 사전 (361 CSV)
# --------------------------------------------------------------------------------------
def load_acupoint_dictionary(csv_path: Path) -> Dict[str, Dict[str, Any]]:
    d: Dict[str, Dict[str, Any]] = {}
    if not csv_path.exists():
        print(f"[Warning] 경혈 CSV 없음: {csv_path}")
        return d
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            raw = (row.get("경혈코드") or "").strip()
            if not raw:
                continue
            sc = std_code(raw)
            m = re.match(r"^([A-Z]+)(\d+)$", sc)
            kr = (row.get("본혈명(한글)") or "").strip()
            cn = (row.get("본혈명(한자)") or "").strip()
            pats = []
            if m:
                # 대소문자 구분: 'Sp5'(삼차신경 척수핵)·'li4' 같은 다른 약어가 경혈 코드로 잡히지 않게 대문자 경락 기호만 인정
                pats.append(re.compile(rf"\b{m.group(1)}[- ]?0?{m.group(2)}\b"))
            if cn:
                # 지명·기관명(商丘市, 太原医院 등) 제외
                pats.append(re.compile(re.escape(cn) + CN_PLACE_LOOKAHEAD))
            if kr:
                # 한글 2자 이름 오탐 방지: 뒤에 '혈' 또는 괄호/코드가 따라오는 경우만 본문 일치로 인정
                pats.append(re.compile(re.escape(kr) + r"(?=혈|\s*[\(（]|\s*[A-Z]{1,2}\d)"))
            d[sc] = {"std_code": sc, "raw_code": raw, "kr": kr, "cn": cn, "meridian_kr": (row.get("소속경락(한글)") or "").strip(),
                     "meridian_cn": (row.get("소속경락(한자)") or "").strip(), "patterns": pats}
    return d


# --------------------------------------------------------------------------------------
# 경혈 언급 판정 (외부 검색 결과의 경혈 연결 전 필수 검사, 2026-10-23)
#   경혈 병음이 지명(太原 Taiyuan, 商丘 Shangqiu)·기관명과 같거나, 코드가 다른 약어(Sp5 = 삼차신경 척수핵)와 같아
#   무관 논문이 연결되던 문제를 막는다. 제목·초록에 경혈 식별자가 '경혈로서' 나와야 연결한다.
# --------------------------------------------------------------------------------------
CN_PLACE_LOOKAHEAD = r"(?![市县縣区區省镇鎮村]|大学|大學|医院|醫院|医学|醫學|人民|中医院|地区|地區)"
_EN_PLACE_AFTER = re.compile(r"^\W{0,3}(city|province|county|district|municipal|university|hospital|medical|people'?s|"
                             r"college|institute|region|prefecture|china\b|first|second|third|fourth|fifth|central|affiliated|"
                             r"traditional|tcm\b|no\.|normal|cancer|maternal|children|shanxi|henan|area|urban|rural|residents?)", re.I)
_ACU_CONTEXT = re.compile(r"acupunct|electroacupunct|acupoint|acu-point|moxibust|needl|\bTEAS\b|\bpoint\b|穴|针刺|針刺|电针|電針|艾灸", re.I)


def code_variants_regex(sc: str) -> Optional[re.Pattern]:
    m = re.match(r"^([A-Z]+)(\d+)$", sc)
    return re.compile(rf"\b{m.group(1)}[- ]?0?{m.group(2)}\b") if m else None


def acupoint_mentioned(text: str, sc: str, pinyins: List[str], names_cn: List[str]) -> Tuple[bool, str]:
    """제목·초록(text)에 경혈 sc가 경혈로서 언급되었는가. (판정, 근거) 반환."""
    if not text:
        return False, "본문 없음"
    rx = code_variants_regex(sc)
    if rx and rx.search(text):
        return True, "code"
    # 혼합 대소문자(Gb21, Lu9, St36)는 침구 맥락이 있을 때만 인정. 'Sp'는 spinal 약어(Sp5 = 삼차신경 척수핵)와 겹쳐 제외
    m = re.match(r"^([A-Z]+)(\d+)$", sc)
    if m and m.group(1) != "SP" and _ACU_CONTEXT.search(text):
        if re.search(rf"\b{m.group(1)[0]}{m.group(1)[1:].lower()}[- ]?0?{m.group(2)}\b", text):
            return True, "code(mixed-case)"
    for cn in [c for c in names_cn if c]:
        if re.search(re.escape(cn) + CN_PLACE_LOOKAHEAD, text):
            return True, "cn"
    if not _ACU_CONTEXT.search(text):
        return False, "침구 맥락 없음"
    for py in [p for p in pinyins if p]:
        for mm in re.finditer(rf"\b{re.escape(py)}\b", text, re.I):
            if not _EN_PLACE_AFTER.match(text[mm.end():mm.end() + 30]):
                return True, "pinyin"
    return False, "식별자 없음(또는 지명·기관명으로만 등장)"


KNOWN_PINYIN = {  # 로컬 PDF 감지용 보조 병음(침구 맥락이 있을 때만). sync-api가 저장한 entry.pinyin도 함께 쓴다
    "ST36": ["Zusanli"], "LI4": ["Hegu", "Hoku"], "SP6": ["Sanyinjiao"], "PC6": ["Neiguan"], "GV20": ["Baihui"],
    "LR3": ["Taichong"], "GB34": ["Yanglingquan"], "BL40": ["Weizhong"], "CV12": ["Zhongwan"], "ST25": ["Tianshu"],
    "CV4": ["Guanyuan"], "GB20": ["Fengchi"], "LI11": ["Quchi"], "KI3": ["Taixi"], "HT7": ["Shenmen"],
}


def detect_acupoints(text: str, filename: str, adict: Dict[str, Dict[str, Any]], pinyin_map: Optional[Dict[str, List[str]]] = None) -> List[str]:
    found = set()
    fl = filename.lower()
    acu_ctx = bool(text) and bool(_ACU_CONTEXT.search(text))
    for sc, info in adict.items():
        if (info["kr"] and info["kr"] in filename) or (info["cn"] and info["cn"] in filename) or sc.lower() in fl:
            found.add(sc)
            continue
        if text:
            for pat in info["patterns"]:
                if pat.search(text):
                    found.add(sc)
                    break
            else:
                if acu_ctx:
                    # 혼합 대소문자 코드(Li 4, St36, St-36) — 침구 맥락이 있을 때만, 'Sp'는 spinal 약어와 겹쳐 제외
                    m = re.match(r"^([A-Z]+)(\d+)$", sc)
                    if m and m.group(1) != "SP" and re.search(rf"\b{m.group(1)[0]}{m.group(1)[1:].lower()}[- .]{{0,2}}0?{m.group(2)}\b(?![.,]?\d)", text):
                        found.add(sc)
                        continue
                    for py in (pinyin_map or {}).get(sc, []) + KNOWN_PINYIN.get(sc, []):
                        mm = re.search(rf"\b{re.escape(py)}\b", text, re.I)
                        if mm and not _EN_PLACE_AFTER.match(text[mm.end():mm.end() + 30]):
                            found.add(sc)
                            break
    return sorted(found)


# --------------------------------------------------------------------------------------
# DB 로드/저장/로그
# --------------------------------------------------------------------------------------
def load_db(ctx: Ctx) -> Dict[str, Any]:
    with open(ctx.json_path, encoding="utf-8") as f:
        db = json.load(f)
    db.setdefault("_metadata", {})
    db.setdefault("papers", [])
    db.setdefault("acupoint_inverted_index", {})
    db.setdefault("file_manifest", {})
    db["_metadata"].setdefault("sync_log", [])
    return db


def backup_file(ctx: Ctx, path: Path) -> Optional[Path]:
    if not path.exists():
        return None
    ctx.backup_dir.mkdir(parents=True, exist_ok=True)
    dst = ctx.backup_dir / f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}"
    shutil.copy2(path, dst)
    # 1세대만 유지: 같은 stem의 이전 백업 삭제
    for old in sorted(ctx.backup_dir.glob(f"{path.stem}_*{path.suffix}"))[:-1]:
        try:
            old.unlink()
        except Exception:
            pass
    return dst


def save_db(ctx: Ctx, db: Dict[str, Any]) -> None:
    if ctx.dry_run:
        print("[dry-run] JSON 저장 생략")
        return
    backup_file(ctx, ctx.json_path)
    fd, tmp = tempfile.mkstemp(dir=str(ctx.ds), suffix=".json.tmp")
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)
    os.replace(tmp, ctx.json_path)
    print(f"[Saved] {ctx.json_path}")


def log_event(ctx: Ctx, db: Dict[str, Any], action: str, **kw) -> Dict[str, Any]:
    ev = {"ts": now_iso(), "trigger": ctx.trigger, "action": action}
    ev.update(kw)
    db["_metadata"].setdefault("sync_log", []).append(ev)
    db["_metadata"]["sync_log"] = db["_metadata"]["sync_log"][-SYNC_LOG_CAP:]
    if not ctx.dry_run:
        ctx.log_dir.mkdir(parents=True, exist_ok=True)
        with open(ctx.log_dir / f"{datetime.now():%Y-%m}.jsonl", "a", encoding="utf-8") as f:
            f.write(json.dumps(ev, ensure_ascii=False) + "\n")
    return ev


def acupoint_entries(db: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    return {k: v for k, v in db.items() if isinstance(v, dict) and "code" in v}


def local_dir(ctx: Ctx, db: Dict[str, Any]) -> Path:
    """로컬 논문 폴더 해석(PC 독립): 환경변수 → 프로젝트 루트 상위(내 드라이브)/학습자료 → DB 기록값 → 기본값."""
    env = os.environ.get("ACUPOINT_PAPERS_DIR")
    if env and Path(env).exists():
        return Path(env)
    cand = ctx.root.parent / "학습자료" / "실험 및 연구 논문"
    if cand.exists():
        return cand
    stored = db["_metadata"].get("source_local_dir")
    if stored and Path(stored).exists():
        return Path(stored)
    return Path(stored or DEFAULT_LOCAL_DIR)


def next_id(db: Dict[str, Any]) -> int:
    return max([as_int(p.get("id"), 0) for p in db["papers"]] + [0]) + 1


def ensure_entry(db: Dict[str, Any], code: str, ctx: Ctx, name_kr="", name_cn="", meridian="") -> Dict[str, Any]:
    sc = std_code(code)
    if sc not in db or not isinstance(db[sc], dict) or "code" not in db[sc]:
        info = ctx.acupoint_dict().get(sc, {})
        db[sc] = {
            "code": sc, "name_kr": name_kr or info.get("kr", ""), "name_cn": name_cn or info.get("cn", ""),
            "meridian": meridian or (f"{info.get('meridian_kr','')} ({info.get('meridian_cn','')})".strip() if info else ""),
            "innervation_anatomy": "", "spinal_segment": "", "reflex_circuit": "", "landmark_papers": [],
            "standard_anatomy_reference": "Gray's Anatomy & Snell's Clinical Neuroanatomy", "clinical_indications_mapped": "",
            "local_paper_ids": [], "local_papers_count": 0, "local_papers_sample": [], "landmark_paper_ids": [],
            "external_paper_ids": [], "external_papers_count": 0, "last_api_sync_at": "",
        }
    e = db[sc]
    if name_kr and not e.get("name_kr"):
        e["name_kr"] = name_kr
    if name_cn and not e.get("name_cn"):
        e["name_cn"] = name_cn
    return e


# --------------------------------------------------------------------------------------
# 파생값 재계산 (역색인 · 경혈 항목)
# --------------------------------------------------------------------------------------
def recompute_derived(db: Dict[str, Any], ctx: Ctx) -> None:
    papers = db["papers"]
    by_id = {as_int(p.get("id")): p for p in papers}
    inv: Dict[str, List[int]] = {}
    for p in papers:
        if p.get("status") in INACTIVE_STATUSES:
            continue
        for c in as_list(p.get("detected_acupoints")):
            inv.setdefault(std_code(c), []).append(as_int(p["id"]))
    for c in inv:
        inv[c] = sorted(set(inv[c]))
    db["acupoint_inverted_index"] = dict(sorted(inv.items()))

    for c in inv:
        ensure_entry(db, c, ctx)

    for code, e in acupoint_entries(db).items():
        ids = inv.get(code, [])
        local_ids = [i for i in ids if by_id[i].get("origin", "local") == "local" and by_id[i].get("status", "active") == "active"]
        ext_ids = [i for i in ids if by_id[i].get("origin", "local") != "local"]
        e["local_paper_ids"] = local_ids
        e["local_papers_count"] = len(local_ids)
        e["local_papers_sample"] = [{"id": i, "filename": by_id[i].get("filename", ""), "year": by_id[i].get("year"),
                                     "doi": by_id[i].get("doi", "")} for i in local_ids[:5]]
        e["external_paper_ids"] = ext_ids
        e["external_papers_count"] = len(ext_ids)
        lm = [i for i in as_list(e.get("landmark_paper_ids")) if i in by_id and by_id[i].get("status") not in INACTIVE_STATUSES]
        e["landmark_paper_ids"] = lm
        e["landmark_papers"] = [paper_to_landmark(by_id[i]) for i in lm]  # 읽기 호환용

    md = db["_metadata"]
    md["counts"] = {
        "papers_total": sum(1 for p in papers if p.get("status") != "superseded"),
        "local_active": sum(1 for p in papers if p.get("origin", "local") == "local" and p.get("status", "active") == "active"),
        "local_missing": sum(1 for p in papers if p.get("origin", "local") == "local" and p.get("status") == "missing"),
        "local_duplicate": sum(1 for p in papers if p.get("origin", "local") == "local" and p.get("status") == "duplicate"),
        "external": sum(1 for p in papers if p.get("origin", "local") != "local" and p.get("status") != "superseded"),
        "acupoints_indexed": len(inv),
    }
    md["total_local_papers"] = md["counts"]["local_active"]
    md["total_acupoints_indexed"] = md["counts"]["acupoints_indexed"]


INACTIVE_STATUSES = ("superseded", "duplicate")


def mark_local_duplicates(db: Dict[str, Any]) -> int:
    """같은 DOI의 로컬 PDF가 여러 개면 가장 작은 ID만 정식(active)으로 두고 나머지는 duplicate로 표시."""
    groups: Dict[str, List[Dict[str, Any]]] = {}
    for p in db["papers"]:
        if p.get("origin", "local") != "local" or p.get("status") in INACTIVE_STATUSES:
            continue
        d = normalize_doi(p.get("doi"))
        if d:
            groups.setdefault(d, []).append(p)
    n = 0
    for d, lst in groups.items():
        if len(lst) < 2:
            continue
        lst.sort(key=lambda x: as_int(x["id"]))
        keep = lst[0]
        for dup in lst[1:]:
            dup["status"] = "duplicate"
            dup["duplicate_of"] = keep["id"]
            dup["updated_at"] = now_iso()
            keep["detected_acupoints"] = sorted(set(as_list(keep.get("detected_acupoints")) + as_list(dup.get("detected_acupoints"))))
            n += 1
    return n


def paper_to_landmark(p: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "id": p.get("id"), "title": p.get("title", ""), "authors": p.get("authors", ""), "journal": p.get("journal", ""),
        "year": p.get("year"), "volume_issue": p.get("volume_issue", ""), "pmid": p.get("pmid", ""), "doi": p.get("doi", ""),
        "evidence_summary": p.get("evidence_summary") or f"{p.get('title','')} [{p.get('journal','')}, {p.get('year','')}]",
        "source_db": p.get("source_db") or p.get("origin", ""),
    }


# --------------------------------------------------------------------------------------
# 텍스트 추출
# --------------------------------------------------------------------------------------
def extract_text(path: Path) -> Tuple[str, bool]:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf" and fitz:
            doc = fitz.open(lp(path))
            txt = "\n".join(doc[i].get_text() for i in range(min(2, len(doc))))
            doc.close()
            return txt, True
        if ext in (".hwp", ".hwpx") and olefile and olefile.isOleFile(lp(path)):
            ole = olefile.OleFileIO(lp(path))
            if ole.exists("PrvText"):
                data = ole.openstream("PrvText").read()
                ole.close()
                return data.decode("utf-16-le", errors="ignore"), True
            ole.close()
    except Exception:
        pass
    return "", False


PINYIN_MAP: Dict[str, List[str]] = {}   # cmd_scan_local이 DB에서 채운다(경혈 코드 → 병음)


def build_local_record(path: Path, base: Path, adict: Dict[str, Dict[str, Any]], pid: int, existing: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    text, ok = extract_text(path)
    meta = parse_filename_meta(path.name)
    doi = pmid = ""
    if text:
        m = re.search(r"\b(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)", text)
        if m:
            doi = normalize_doi(m.group(1))
        m = re.search(r"\bPMID[:\s]*(\d{7,9})\b", text, re.IGNORECASE)
        if m:
            pmid = m.group(1)
    snippet = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", re.sub(r"\s+", " ", text)).strip()[:300] if text else ""
    st = os.stat(lp(path))
    rec = dict(existing or {})
    rec.update({
        "id": pid, "origin": "local", "status": "active",
        "title": meta["title"], "authors": meta["author"], "journal": meta["journal"], "year": meta["year"],
        "volume_issue": str(meta["year"] or ""), "doi": doi or rec.get("doi", ""), "pmid": pmid or rec.get("pmid", ""),
        "abstract_snippet": snippet, "snippet": snippet,
        # 본문 감지 ∪ sync-api가 초록으로 확인해 붙인 연결(api_linked_acupoints) — 전면 재스캔에도 API 연결이 사라지지 않게 한다
        "detected_acupoints": sorted(set(detect_acupoints(text, path.name, adict, PINYIN_MAP)) | set(as_list(rec.get("api_linked_acupoints"))) | set(as_list(rec.get("manual_linked_acupoints")))),
        "keywords": rec.get("keywords", []), "mesh": rec.get("mesh", []), "language": rec.get("language", ""),
        "evidence_level": rec.get("evidence_level"),
        "file_path": relpath_norm(path, base), "filename": path.name, "ext": path.suffix.lower(),
        "file_size": st.st_size, "mtime": round(st.st_mtime, 3), "content_hash": file_hash(path),
        "parsed_filename": meta, "text_extracted": ok,
        "added_at": rec.get("added_at") or now_iso(), "updated_at": now_iso(),
    })
    return rec


# --------------------------------------------------------------------------------------
# scan-local
# --------------------------------------------------------------------------------------
def cmd_scan_local(ctx: Ctx, db: Dict[str, Any], full_rescan=False, enrich=False, enrich_limit: int = 0) -> Dict[str, Any]:
    PINYIN_MAP.clear()
    PINYIN_MAP.update(_pinyin_map_from_logs(db))
    for _sc, _e in acupoint_entries(db).items():
        if _e.get("pinyin") and _e["pinyin"] not in PINYIN_MAP.setdefault(_sc, []):
            PINYIN_MAP[_sc].append(_e["pinyin"])
    t0 = time.time()
    base = local_dir(ctx, db)
    adict = ctx.acupoint_dict()
    print(f"\n[scan-local] 대상 폴더: {base} | full_rescan={full_rescan} enrich={enrich}")
    if not base.exists():
        print("[Error] 로컬 폴더가 없습니다.")
        log_event(ctx, db, "scan-local", source="local", http_status=None, note="folder missing")
        return {"error": "folder missing"}

    disk: Dict[str, Path] = {}
    for root, _, files in os.walk(base):
        for fn in files:
            p = Path(root) / fn
            if p.suffix.lower() in LOCAL_EXTS:
                disk[relpath_norm(p, base)] = p

    manifest: Dict[str, Dict[str, Any]] = db.get("file_manifest", {})
    by_id = {as_int(p.get("id")): p for p in db["papers"]}
    hash_to_rel = {m.get("hash"): rel for rel, m in manifest.items() if m.get("hash")}

    added = updated = moved = missing = unchanged = 0
    new_paths: List[str] = []
    seen_rel = set()

    for rel in sorted(disk):
        p = disk[rel]
        st = os.stat(lp(p))
        prev = manifest.get(rel)
        if prev and prev.get("id") in by_id:
            seen_rel.add(rel)
            rec = by_id[prev["id"]]
            same_stat = (prev.get("size") == st.st_size and abs(float(prev.get("mtime", 0)) - st.st_mtime) < 1.0)
            if same_stat and not full_rescan:
                if rec.get("status") == "missing":  # 유실 복귀만 active로 되돌림 (duplicate 표시는 유지)
                    rec["status"] = "active"
                unchanged += 1
                continue
            h = file_hash(p)
            if h == prev.get("hash") and not full_rescan:
                prev["mtime"] = round(st.st_mtime, 3)
                rec["mtime"] = prev["mtime"]
                if rec.get("status") == "missing":
                    rec["status"] = "active"
                unchanged += 1
                continue
            newrec = build_local_record(p, base, adict, prev["id"], existing=rec)
            by_id[prev["id"]].clear()
            by_id[prev["id"]].update(newrec)
            manifest[rel] = {"id": prev["id"], "size": st.st_size, "mtime": round(st.st_mtime, 3), "hash": newrec["content_hash"]}
            updated += 1
            continue
        # 경로가 manifest에 없음 → 이동 여부(동일 해시) 확인
        h = file_hash(p)
        old_rel = hash_to_rel.get(h)
        if old_rel and old_rel not in disk and old_rel in manifest and manifest[old_rel].get("id") in by_id:
            pid = manifest[old_rel]["id"]
            rec = by_id[pid]
            rec.update({"file_path": rel, "filename": p.name, "status": "active", "mtime": round(st.st_mtime, 3), "updated_at": now_iso()})
            manifest[rel] = dict(manifest.pop(old_rel), mtime=round(st.st_mtime, 3))
            seen_rel.add(rel)
            moved += 1
            continue
        new_paths.append(rel)

    # 삭제 감지
    for rel in list(manifest):
        if rel not in disk and rel not in seen_rel:
            pid = manifest[rel].get("id")
            if pid in by_id and by_id[pid].get("status") == "active":
                by_id[pid]["status"] = "missing"
                by_id[pid]["updated_at"] = now_iso()
                missing += 1

    # 신규 파일: 경로 정렬 후 순차 ID
    nid = next_id(db)
    for rel in sorted(new_paths):
        p = disk[rel]
        rec = build_local_record(p, base, adict, nid)
        db["papers"].append(rec)
        by_id[nid] = rec
        manifest[rel] = {"id": nid, "size": rec["file_size"], "mtime": rec["mtime"], "hash": rec["content_hash"]}
        nid += 1
        added += 1
        if added % 25 == 0:
            print(f"    신규 처리 {added}/{len(new_paths)}")

    db["file_manifest"] = dict(sorted(manifest.items()))
    dups = mark_local_duplicates(db) + mark_local_title_duplicates(db)

    enriched = 0
    if enrich:
        enriched = enrich_local_by_doi(ctx, db, enrich_limit)

    recompute_derived(db, ctx)
    dur = int((time.time() - t0) * 1000)
    db["_metadata"]["last_local_scan"] = {"at": now_iso(), "files_seen": len(disk), "added": added, "updated": updated,
                                          "moved": moved, "missing": missing, "unchanged": unchanged, "duplicates_marked": dups,
                                          "enriched": enriched, "duration_ms": dur}
    log_event(ctx, db, "scan-local", source="local", http_status=None, fetched=len(disk), added=added, deduped=dups,
              duration_ms=dur, note=f"updated={updated} moved={moved} missing={missing} enriched={enriched}")
    print(f"[scan-local] files={len(disk)} added={added} updated={updated} moved={moved} missing={missing} unchanged={unchanged} dup_marked={dups} enriched={enriched} ({dur} ms)")
    return db["_metadata"]["last_local_scan"]


def enrich_local_by_doi(ctx: Ctx, db: Dict[str, Any], limit: int = 0) -> int:
    """DOI 보유 로컬 레코드에 PubMed(efetch) 메타(PMID·제목·초록·MeSH)를 보강. limit>0이면 상위 N편만."""
    if not requests:
        return 0
    targets = [p for p in db["papers"] if p.get("origin", "local") == "local" and p.get("doi") and not p.get("pmid") and not p.get("enrich_checked_at")]
    if limit > 0:
        targets = targets[:limit]
    print(f"[enrich] 대상 {len(targets)}편 (DOI 보유 · PMID 미확인{', limit=' + str(limit) if limit else ''})")
    n = 0
    for i, p in enumerate(targets, 1):
        ids, status = pubmed_esearch(ctx, f'{p["doi"]}[AID]', retmax=1)
        p["enrich_checked_at"] = now_iso()
        if status == 200 and ids:
            recs, st2 = pubmed_efetch(ctx, ids)
            if recs:
                merge_into(p, recs[0], source="pubmed")
                n += 1
        if i % 50 == 0:
            print(f"    enrich {i}/{len(targets)} (보강 {n})")
    log_event(ctx, db, "enrich", source="pubmed", http_status=200, fetched=len(targets), added=n)
    return n


# --------------------------------------------------------------------------------------
# 외부 API 수집기 (각각 (records, status) 반환)
# --------------------------------------------------------------------------------------
def http_get(url: str, params: Optional[dict] = None, retries: int = 1) -> Tuple[Optional["requests.Response"], Optional[int], str]:
    if not requests:
        return None, None, "requests 미설치"
    last_err = ""
    for attempt in range(retries + 1):
        try:
            r = requests.get(url, params=params, headers={"User-Agent": UA}, timeout=HTTP_TIMEOUT)
            if r.status_code == 200:
                return r, 200, ""
            last_err = f"HTTP {r.status_code}"
            if r.status_code in (429, 500, 502, 503) and attempt < retries:
                time.sleep(1.5)
                continue
            return r, r.status_code, last_err
        except Exception as e:
            last_err = type(e).__name__
            if attempt < retries:
                time.sleep(1.0)
    return None, None, last_err


def pubmed_esearch(ctx: Ctx, term: str, retmax: int = 8, sort: str = "relevance") -> Tuple[List[str], Optional[int]]:
    params = {"db": "pubmed", "term": term, "retmode": "json", "retmax": retmax, "sort": sort}
    if NCBI_KEY:
        params["api_key"] = NCBI_KEY
    r, status, _ = http_get("https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi", params)
    time.sleep(NCBI_SLEEP)
    if status != 200:
        return [], status
    return r.json().get("esearchresult", {}).get("idlist", []), 200


def pubmed_efetch(ctx: Ctx, pmids: List[str]) -> Tuple[List[Dict[str, Any]], Optional[int]]:
    if not pmids:
        return [], 200
    params = {"db": "pubmed", "id": ",".join(pmids), "retmode": "xml"}
    if NCBI_KEY:
        params["api_key"] = NCBI_KEY
    r, status, _ = http_get("https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi", params)
    time.sleep(NCBI_SLEEP)
    if status != 200:
        return [], status
    out = []
    try:
        root = ET.fromstring(r.content)
    except ET.ParseError:
        return [], status
    for art in root.findall(".//PubmedArticle"):
        mc = art.find("MedlineCitation")
        a = mc.find("Article") if mc is not None else None
        if a is None:
            continue
        pmid = mc.findtext("PMID", "")
        title = re.sub(r"</?[bi]>", "", "".join(a.find("ArticleTitle").itertext()) if a.find("ArticleTitle") is not None else "").strip()
        title = title.strip("[]").rstrip(".") if title.startswith("[") else title
        abst_parts = []
        for x in a.findall("Abstract/AbstractText"):
            lab = x.get("Label")
            txt = "".join(x.itertext()).strip()
            abst_parts.append(f"{lab}: {txt}" if lab else txt)
        abstract = " ".join(abst_parts)
        journal = a.findtext("Journal/ISOAbbreviation") or a.findtext("Journal/Title") or ""
        year = a.findtext("Journal/JournalIssue/PubDate/Year") or ""
        if not year:
            md_ = a.findtext("Journal/JournalIssue/PubDate/MedlineDate") or ""
            m = re.search(r"(19|20)\d\d", md_)
            year = m.group(0) if m else ""
        vol = a.findtext("Journal/JournalIssue/Volume") or ""
        iss = a.findtext("Journal/JournalIssue/Issue") or ""
        pages = a.findtext("Pagination/MedlinePgn") or ""
        vol_str = f"{vol}({iss}): {pages}".strip() if vol else pages
        authors = []
        for au in a.findall("AuthorList/Author"):
            ln, ini = au.findtext("LastName") or "", au.findtext("Initials") or ""
            if ln:
                authors.append(f"{ln} {ini}".strip())
        doi = ""
        # 논문 자신의 식별자만 본다(PubmedData/ArticleIdList). './/ArticleIdList'는 참고문헌 목록의 DOI까지 포함하므로 쓰지 않는다.
        for e in art.findall("PubmedData/ArticleIdList/ArticleId"):
            if e.get("IdType") == "doi":
                doi = normalize_doi(e.text)
                break
        if not doi:  # ArticleIdList에 없으면 ELocationID(doi)에서 보완
            for e in a.findall("ELocationID"):
                if e.get("EIdType") == "doi":
                    doi = normalize_doi(e.text)
                    break
        out.append({
            "title": title, "authors": ", ".join(authors[:6]) + (" et al." if len(authors) > 6 else ""), "journal": journal,
            "year": as_int(year, None), "volume_issue": vol_str, "pmid": pmid, "doi": doi,
            "abstract_snippet": abstract[:300], "abstract_full_len": len(abstract), "_fulltext": f"{title} {abstract}",
            "keywords": [k.text.strip() for k in mc.findall("KeywordList/Keyword") if k.text],
            "mesh": [m.findtext("DescriptorName") for m in mc.findall("MeshHeadingList/MeshHeading") if m.findtext("DescriptorName")],
            "language": a.findtext("Language") or "", "source_db": "PubMed",
            "url": f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/",
        })
    return out, 200


def search_pubmed(ctx: Ctx, code: str, pinyin: str, name_cn: str, max_n: int) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    # [tiab]: 저자 소속(예: 太原 Taiyuan 소재 병원)까지 검색되던 것을 제목·초록으로 한정 (2026-10-23)
    terms = " OR ".join(f'"{t}"[tiab]' for t in [code, pinyin, name_cn] if t)
    queries = [
        (f"({terms}) AND (acupuncture OR electroacupuncture) AND (nerve OR reflex OR autonomic OR mechanism OR anatomy)", "relevance"),
        (f'({terms}) AND ("Zhongguo Zhen Jiu"[Journal] OR "Zhen Ci Yan Jiu"[Journal])', "pub_date"),
    ]
    recs, statuses, ids_all = [], [], []
    for q, sort in queries:
        ids, st = pubmed_esearch(ctx, q, retmax=max_n, sort=sort)
        statuses.append(st)
        ids_all.extend(i for i in ids if i not in ids_all)
    if ids_all:
        got, st = pubmed_efetch(ctx, ids_all)
        statuses.append(st)
        for g in got:
            g["source_query"] = queries[0][0] if g.get("journal", "").lower() not in ("zhongguo zhen jiu", "zhen ci yan jiu") else queries[1][0]
        recs = got
    ok = all(s == 200 for s in statuses) and bool(statuses)
    return recs, {"source": "pubmed", "http_status": 200 if ok else (next((s for s in statuses if s != 200), None)), "fetched": len(recs), "queries": [q for q, _ in queries]}


def openalex_abstract(inv: Optional[dict]) -> str:
    if not inv:
        return ""
    pos = {}
    for w, ps in inv.items():
        for p in ps:
            pos[p] = w
    return " ".join(pos[i] for i in sorted(pos))


def openalex_record(w: dict, query: str) -> Dict[str, Any]:
    src = ((w.get("primary_location") or {}).get("source") or {})
    ids = w.get("ids") or {}
    pm = re.search(r"(\d+)$", ids.get("pmid", "") or "")
    bib = w.get("biblio") or {}
    vol_str = f"{bib.get('volume') or ''}({bib.get('issue') or ''}): {bib.get('first_page') or ''}-{bib.get('last_page') or ''}".replace("(): -", "").strip()
    authors = [((a.get("author") or {}).get("display_name") or "") for a in (w.get("authorships") or [])]
    authors = [a for a in authors if a]
    abstract = openalex_abstract(w.get("abstract_inverted_index"))
    return {
        "title": (w.get("title") or "").strip("[]"), "authors": ", ".join(authors[:6]) + (" et al." if len(authors) > 6 else ""),
        "journal": src.get("display_name") or "", "year": w.get("publication_year"), "volume_issue": vol_str,
        "pmid": pm.group(1) if pm else "", "doi": normalize_doi(w.get("doi")), "abstract_snippet": abstract[:300],
        "keywords": [k.get("display_name") for k in (w.get("keywords") or []) if k.get("display_name")][:10], "mesh": [],
        "language": w.get("language") or "", "source_db": "OpenAlex", "source_query": query, "url": w.get("id") or "",
        "_fulltext": f"{(w.get('title') or '')} {abstract}",
    }


def search_openalex(ctx: Ctx, code: str, pinyin: str, max_n: int) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    kw = " OR ".join(f'"{t}"' for t in [pinyin, code] if t) or code
    sel = "id,doi,title,publication_year,ids,primary_location,authorships,abstract_inverted_index,biblio,language,keywords"
    calls = [(f"doi_starts_with:{p},title_and_abstract.search:{kw}", None) for p in CN_JOURNAL_PREFIXES]
    calls.append((f"has_doi:true,title_and_abstract.search:{kw} AND (acupuncture OR electroacupuncture)", None))
    recs, statuses = [], []
    for flt, _ in calls:
        params = {"filter": flt, "per-page": max_n, "sort": "publication_year:desc", "select": sel}
        if "mailto:" in UA and "local.workspace" not in UA:
            params["mailto"] = UA.split("mailto:")[1].rstrip(")")
        r, st, err = http_get("https://api.openalex.org/works", params)
        statuses.append(st)
        if st != 200:
            print(f"  [OpenAlex] {err} filter={flt[:60]}")
            continue
        for w in (r.json().get("results") or []):
            if w.get("title"):
                recs.append(openalex_record(w, flt))
    ok = all(s == 200 for s in statuses)
    return recs, {"source": "openalex", "http_status": 200 if ok else next((s for s in statuses if s != 200), None), "fetched": len(recs), "queries": [c[0] for c in calls]}


def search_crossref(ctx: Ctx, code: str, pinyin: str, max_n: int) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    q = f"{pinyin or code} acupuncture nerve reflex"
    params = {"query.bibliographic": q, "rows": max_n, "sort": "relevance", "filter": "type:journal-article"}
    r, st, err = http_get("https://api.crossref.org/works", params)
    recs = []
    # Crossref query.bibliographic 은 느슨한 OR 매칭이므로 제목·초록에 경혈 식별자가 실제로 있어야 채택
    needles = [t.lower() for t in [code, pinyin] if t]
    m_ = re.match(r"^([A-Z]+)(\d+)$", code)
    if m_:
        needles.append(f"{m_.group(1).lower()} {m_.group(2)}")
    if st == 200:
        for it in (r.json().get("message", {}).get("items") or []):
            title = (it.get("title") or [""])[0]
            if not title:
                continue
            hay = f"{title} {it.get('abstract') or ''}".lower()
            if not any(n in hay for n in needles):
                continue
            issued = ((it.get("issued") or {}).get("date-parts") or [[None]])[0]
            authors = [f"{a.get('given','')} {a.get('family','')}".strip() for a in (it.get("author") or [])]
            recs.append({
                "title": title, "authors": ", ".join([a for a in authors if a][:6]), "journal": (it.get("container-title") or [""])[0],
                "year": issued[0] if issued else None, "volume_issue": f"{it.get('volume','')}({it.get('issue','')})".replace("()", ""),
                "pmid": "", "doi": normalize_doi(it.get("DOI")), "abstract_snippet": re.sub(r"<[^>]+>", "", it.get("abstract") or "")[:300],
                "keywords": [], "mesh": [], "language": it.get("language") or "", "source_db": "Crossref", "source_query": q,
                "url": it.get("URL") or "",
            })
    else:
        print(f"  [Crossref] {err}")
    return recs, {"source": "crossref", "http_status": st, "fetched": len(recs), "queries": [q]}


def search_europepmc(ctx: Ctx, code: str, pinyin: str, max_n: int) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    q = f'({" OR ".join(t for t in [code, pinyin] if t)}) AND (acupuncture OR electroacupuncture) AND (nerve OR reflex OR autonomic)'
    params = {"query": q, "format": "json", "resultType": "core", "pageSize": max_n}
    r, st, err = http_get("https://www.ebi.ac.uk/europepmc/webservices/rest/search", params)
    recs = []
    if st == 200:
        for it in (r.json().get("resultList", {}).get("result") or []):
            recs.append({
                "title": (it.get("title") or "").strip("[]"), "authors": it.get("authorString") or "",
                "journal": ((it.get("journalInfo") or {}).get("journal") or {}).get("title") or it.get("journalTitle") or "",
                "year": as_int(it.get("pubYear"), None), "volume_issue": f"{it.get('journalVolume','')}({it.get('issue','')}): {it.get('pageInfo','')}",
                "pmid": it.get("pmid") or "", "doi": normalize_doi(it.get("doi")), "abstract_snippet": (it.get("abstractText") or "")[:300],
                "keywords": [], "mesh": [], "language": it.get("language") or "", "source_db": "EuropePMC", "source_query": q,
                "url": f"https://europepmc.org/abstract/MED/{it.get('pmid')}" if it.get("pmid") else "",
            })
    else:
        print(f"  [EuropePMC] {err}")
    return recs, {"source": "europepmc", "http_status": st, "fetched": len(recs), "queries": [q]}


# --------------------------------------------------------------------------------------
# 레지스트리 중복 제거 · 병합
# --------------------------------------------------------------------------------------
def build_registry_index(db: Dict[str, Any]) -> Tuple[Dict[str, int], Dict[str, int], Dict[str, int]]:
    by_doi, by_pmid, by_hash = {}, {}, {}
    for p in db["papers"]:
        if p.get("status") in INACTIVE_STATUSES:
            continue
        pid = as_int(p["id"])
        d = normalize_doi(p.get("doi"))
        if d:
            by_doi.setdefault(d, pid)
        pm = str(p.get("pmid") or "").strip()
        if pm and pm != "None":
            by_pmid.setdefault(pm, pid)
        h = title_hash(p)
        if h:
            by_hash.setdefault(h, pid)
    return by_doi, by_pmid, by_hash


def find_duplicate(c: Dict[str, Any], by_doi, by_pmid, by_hash) -> Optional[int]:
    d = normalize_doi(c.get("doi"))
    if d and d in by_doi:
        return by_doi[d]
    pm = str(c.get("pmid") or "").strip()
    if pm and pm in by_pmid:
        return by_pmid[pm]
    h = title_hash(c)
    if h and h in by_hash:
        return by_hash[h]
    return None


def merge_into(target: Dict[str, Any], cand: Dict[str, Any], source: str) -> None:
    for k in ("pmid", "doi", "journal", "volume_issue", "language", "url"):
        if not target.get(k) and cand.get(k):
            target[k] = cand[k]
    if cand.get("year") and not target.get("year"):
        target["year"] = cand["year"]
    if target.get("origin", "local") == "local":
        # 로컬 파일명 기반 제목보다 공식 제목이 정확하면 보존하되 원래 제목은 유지
        if cand.get("title") and not target.get("official_title"):
            target["official_title"] = cand["title"]
    if cand.get("abstract_snippet") and len(cand["abstract_snippet"]) > len(target.get("abstract_snippet") or ""):
        target["abstract_snippet"] = cand["abstract_snippet"]
    for k in ("keywords", "mesh"):
        if cand.get(k):
            target[k] = sorted(set(as_list(target.get(k)) + list(cand[k])))
    afi = set(as_list(target.get("also_found_in")))
    afi.add(source)
    target["also_found_in"] = sorted(afi)
    target["updated_at"] = now_iso()


def cmd_sync_api(ctx: Ctx, db: Dict[str, Any], code: str, name_kr: str, name_cn: str, pinyin: str, meridian: str, max_n: int) -> Dict[str, Any]:
    t0 = time.time()
    sc = std_code(code)
    adict = ctx.acupoint_dict()
    info = adict.get(sc, {})
    name_kr = name_kr or info.get("kr", "")
    name_cn = name_cn or info.get("cn", "")
    entry = ensure_entry(db, sc, ctx, name_kr, name_cn, meridian)
    if pinyin:
        entry["pinyin"] = pinyin
    pinyins = [x for x in {pinyin, entry.get("pinyin", "")} if x]
    names_cn = [x for x in {name_cn, info.get("cn", ""), entry.get("name_cn", "")} if x]
    rejected: List[Dict[str, Any]] = []
    print(f"\n[sync-api] {sc} ({name_kr} / {name_cn} / {pinyin}) max_per_source={max_n}")

    results: List[Tuple[List[Dict[str, Any]], Dict[str, Any]]] = []
    pm_recs, pm_st = search_pubmed(ctx, sc, pinyin, name_cn, max_n)
    results.append((pm_recs, pm_st))
    print(f"  PubMed    : status={pm_st['http_status']} fetched={len(pm_recs)}")
    if pm_st["http_status"] != 200:
        ep_recs, ep_st = search_europepmc(ctx, sc, pinyin, max_n)
        results.append((ep_recs, ep_st))
        print(f"  EuropePMC : status={ep_st['http_status']} fetched={len(ep_recs)} (PubMed 폴백)")
    oa_recs, oa_st = search_openalex(ctx, sc, pinyin, max_n)
    results.append((oa_recs, oa_st))
    print(f"  OpenAlex  : status={oa_st['http_status']} fetched={len(oa_recs)}")
    cr_recs, cr_st = search_crossref(ctx, sc, pinyin, max_n)
    results.append((cr_recs, cr_st))
    print(f"  Crossref  : status={cr_st['http_status']} fetched={len(cr_recs)}")

    by_doi, by_pmid, by_hash = build_registry_index(db)
    by_id = {as_int(p["id"]): p for p in db["papers"]}
    nid = next_id(db)
    added = merged = deduped_batch = 0
    touched_ids: List[int] = []
    for recs, st in results:
        src = st["source"]
        for c in recs:
            ftxt = c.pop("_fulltext", "") or f"{c.get('title', '')} {c.get('abstract_snippet', '')}"
            ok, why = acupoint_mentioned(ftxt, sc, pinyins, names_cn)
            if not ok:
                rejected.append({"source": src, "pmid": c.get("pmid", ""), "doi": c.get("doi", ""), "title": (c.get("title") or "")[:120], "reason": why})
                continue
            dup = find_duplicate(c, by_doi, by_pmid, by_hash)
            if dup is not None:
                tgt = by_id[dup]
                # 교차 병합: DOI로 찾은 레코드와 PMID로 찾은 레코드가 다르면(예: PubMed가 DOI 없이 먼저 생성)
                # 이번 배치에서 생성된 쪽을 superseded로 돌리고 PMID를 정식 레코드로 모은다
                cpm = str(c.get("pmid") or "").strip()
                other = by_pmid.get(cpm) if cpm else None
                if other is not None and other != dup:
                    o = by_id[other]
                    if o.get("fetched_at", "") and o.get("origin") != "local" and other in touched_ids:
                        o["status"] = "superseded"
                        o["superseded_reason"] = f"same PMID {cpm} as id {dup} (cross-source merge)"
                        o["updated_at"] = now_iso()
                        merge_into(tgt, o, source=o.get("origin", src))
                        tgt["detected_acupoints"] = sorted(set(as_list(tgt.get("detected_acupoints")) + as_list(o.get("detected_acupoints"))))
                        touched_ids.remove(other)
                        added -= 1
                        deduped_batch += 1
                    by_pmid[cpm] = dup
                if dup in touched_ids:
                    deduped_batch += 1
                else:
                    merge_into(tgt, c, source=src)
                    merged += 1
                if cpm and cpm not in by_pmid:
                    by_pmid[cpm] = dup
                if sc not in [std_code(x) for x in as_list(tgt.get("detected_acupoints"))]:
                    tgt["detected_acupoints"] = sorted(set(as_list(tgt.get("detected_acupoints")) + [sc]))
                tgt["api_linked_acupoints"] = sorted(set(as_list(tgt.get("api_linked_acupoints")) + [sc]))
                if dup not in touched_ids:
                    touched_ids.append(dup)
                continue
            text = f"{c.get('title','')} {c.get('abstract_snippet','')}"
            det = set(detect_acupoints(text, "", adict))
            det.add(sc)
            rec = {
                "id": nid, "origin": ORIGIN_MAP.get(src, src), "status": "active",
                "title": c.get("title", ""), "authors": c.get("authors", ""), "journal": c.get("journal", ""), "year": c.get("year"),
                "volume_issue": c.get("volume_issue", ""), "doi": normalize_doi(c.get("doi")), "pmid": str(c.get("pmid") or ""),
                "abstract_snippet": c.get("abstract_snippet", ""), "snippet": c.get("abstract_snippet", ""),
                "detected_acupoints": sorted(det), "keywords": c.get("keywords", []), "mesh": c.get("mesh", []),
                "language": c.get("language", ""), "evidence_level": None,
                "source_db": c.get("source_db", src), "source_query": c.get("source_query", ""), "fetched_at": now_iso(),
                "url": c.get("url", ""), "cn_title": "", "cn_abstract": "", "also_found_in": [src],
                "filename": "", "file_path": "", "parsed_filename": {},
                "evidence_summary": f"{c.get('title','')} [{c.get('journal','')}, {c.get('year','')}]",
                "added_at": now_iso(), "updated_at": now_iso(),
            }
            db["papers"].append(rec)
            by_id[nid] = rec
            if rec["doi"]:
                by_doi[rec["doi"]] = nid
            if rec["pmid"]:
                by_pmid[rec["pmid"]] = nid
            h = title_hash(rec)
            if h:
                by_hash[h] = nid
            touched_ids.append(nid)
            nid += 1
            added += 1

    lm = as_list(entry.get("landmark_paper_ids"))
    for i in touched_ids:
        if i not in lm:
            lm.append(i)
    entry["landmark_paper_ids"] = lm
    entry["last_api_sync_at"] = now_iso()

    recompute_derived(db, ctx)
    dur = int((time.time() - t0) * 1000)
    src_status = {st["source"]: ("ok" if st["http_status"] == 200 else f"fail({st['http_status']})") for _, st in results}
    db["_metadata"]["last_api_sync"] = {"at": now_iso(), "acupoint": sc, "sources": src_status}
    for recs, st in results:
        log_event(ctx, db, "sync-api", source=st["source"], acupoint=sc, query=" | ".join(st.get("queries", []))[:400],
                  http_status=st["http_status"], fetched=st["fetched"], duration_ms=dur)
    log_event(ctx, db, "sync-api-summary", source="all", acupoint=sc, http_status=None,
              fetched=sum(len(r) for r, _ in results), deduped=merged + deduped_batch, added=added, duration_ms=dur,
              note=f"merged_into_existing={merged} batch_dups={deduped_batch} rejected_not_mentioned={len(rejected)} sources={src_status}")
    if rejected:
        md = db["_metadata"].setdefault("relevance_rejections", {})
        md[sc] = (md.get(sc, []) + [{**r, "ts": now_iso()} for r in rejected])[-200:]
    print(f"[sync-api] 후보 {sum(len(r) for r, _ in results)} → 신규 {added}, 기존 병합 {merged}, 배치 내 중복 {deduped_batch}, "
          f"경혈 미언급 제외 {len(rejected)} | {sc} 랜드마크 {len(lm)}편 ({dur} ms)")
    all_failed = all(st["http_status"] != 200 for _, st in results)
    return {"code": sc, "added": added, "merged": merged, "rejected": len(rejected), "sources": src_status, "all_failed": all_failed}


# --------------------------------------------------------------------------------------
# relink — 기존 외부 논문의 경혈 연결 재검사 (2026-10-23)
# --------------------------------------------------------------------------------------
def _pinyin_map_from_logs(db: Dict[str, Any]) -> Dict[str, List[str]]:
    """sync_log의 PubMed 쿼리 ("CODE" OR "Pinyin" OR "汉字")에서 경혈별 병음을 복원한다."""
    out: Dict[str, set] = {}
    for ev in db["_metadata"].get("sync_log", []):
        q = ev.get("query") or ""
        sc = std_code(ev.get("acupoint") or "")
        if not sc or ev.get("source") != "pubmed":
            continue
        m = re.search(r'\(([^()]*?)\)\s*AND', q)
        if not m:
            continue
        for t in re.findall(r'"([^"]+)"', m.group(1)):
            if re.fullmatch(r"[A-Za-z][a-z]+", t):
                out.setdefault(sc, set()).add(t)
    return {k: sorted(v) for k, v in out.items()}


def _refetch_texts(ctx: Ctx, papers: List[Dict[str, Any]]) -> Dict[int, str]:
    """외부 논문의 제목+초록 전문을 PMID(PubMed efetch) 또는 DOI(OpenAlex)로 다시 받는다."""
    texts: Dict[int, str] = {}
    by_pmid = {str(p["pmid"]): p for p in papers if str(p.get("pmid") or "").strip() not in ("", "None")}
    pm = list(by_pmid)
    for i in range(0, len(pm), 150):
        got, st = pubmed_efetch(ctx, pm[i:i + 150])
        for g in got:
            p = by_pmid.get(str(g.get("pmid")))
            if p is not None:
                texts[as_int(p["id"])] = g.get("_fulltext", "")
    rest = [p for p in papers if as_int(p["id"]) not in texts and normalize_doi(p.get("doi"))]
    for i in range(0, len(rest), 40):
        chunk = rest[i:i + 40]
        flt = "doi:" + "|".join(normalize_doi(p["doi"]) for p in chunk)
        r, st, err = http_get("https://api.openalex.org/works", {"filter": flt, "per-page": 50, "select": "doi,title,abstract_inverted_index"})
        if st != 200:
            continue
        idx = {normalize_doi(p["doi"]): p for p in chunk}
        for w in (r.json().get("results") or []):
            p = idx.get(normalize_doi(w.get("doi")))
            if p is not None:
                texts[as_int(p["id"])] = f"{w.get('title') or ''} {openalex_abstract(w.get('abstract_inverted_index'))}"
    return texts


def cmd_relink(ctx: Ctx, db: Dict[str, Any], refetch: bool = True) -> Dict[str, Any]:
    t0 = time.time()
    adict = ctx.acupoint_dict()
    pmap = _pinyin_map_from_logs(db)
    for sc, e in acupoint_entries(db).items():
        if e.get("pinyin"):
            pmap.setdefault(sc, [])
            if e["pinyin"] not in pmap[sc]:
                pmap[sc].append(e["pinyin"])
    # manual(사람이 직접 등록·연결한 레코드)은 재검사 대상에서 뺀다
    ext = [p for p in db["papers"] if p.get("origin", "local") not in ("local", "manual") and p.get("status") not in INACTIVE_STATUSES]
    texts = _refetch_texts(ctx, ext) if refetch else {}
    removed_total, kept_total, changed = 0, 0, 0
    per_code: Dict[str, int] = {}
    for p in ext:
        pid = as_int(p["id"])
        txt = texts.get(pid) or f"{p.get('title', '')} {p.get('official_title', '')} {p.get('abstract_snippet', '')}"
        full = pid in texts
        keep, drop = [], []
        for c in as_list(p.get("detected_acupoints")):
            sc = std_code(c)
            info = adict.get(sc, {})
            names_cn = [x for x in {info.get("cn", ""), (db.get(sc) or {}).get("name_cn", "")} if x]
            ok, why = acupoint_mentioned(txt, sc, pmap.get(sc, []), names_cn)
            # 보고서에서 사람이 확인한 연결(manual_linked_acupoints)은 유지. 초록 전문을 못 받은 레코드는 확실한 경우만 제거
            if ok or sc in [std_code(x) for x in as_list(p.get("manual_linked_acupoints"))] or (not full and why == "본문 없음"):
                keep.append(sc)
            else:
                drop.append({"code": sc, "reason": why, "fulltext": full})
        if drop:
            p["detected_acupoints"] = sorted(set(keep))
            p.setdefault("relink_removed", []).extend([{**d, "ts": now_iso()} for d in drop])
            p["updated_at"] = now_iso()
            changed += 1
            removed_total += len(drop)
            for d in drop:
                per_code[d["code"]] = per_code.get(d["code"], 0) + 1
        kept_total += len(keep)
    det_by_id = {as_int(q["id"]): {std_code(x) for x in as_list(q.get("detected_acupoints"))} for q in db["papers"]}
    for sc, e in acupoint_entries(db).items():   # 연결이 끊긴 논문은 랜드마크 목록에서도 뺀다
        e["landmark_paper_ids"] = [i for i in as_list(e.get("landmark_paper_ids")) if sc in det_by_id.get(as_int(i), set())]
    recompute_derived(db, ctx)
    dur = int((time.time() - t0) * 1000)
    log_event(ctx, db, "relink", source="all", acupoint="", http_status=None, fetched=len(texts), deduped=None, added=None, duration_ms=dur,
              note=f"external_checked={len(ext)} fulltext={len(texts)} papers_changed={changed} links_removed={removed_total}")
    top = sorted(per_code.items(), key=lambda x: -x[1])[:15]
    print(f"[relink] 외부 논문 {len(ext)}편 검사(초록 재수집 {len(texts)}편) → 연결 수정 {changed}편, 제거 {removed_total}건, 유지 {kept_total}건 ({dur} ms)")
    print(f"         경혈별 제거 상위: {top}")
    return {"changed": changed, "removed": removed_total, "per_code": per_code}


def mark_local_title_duplicates(db: Dict[str, Any]) -> int:
    """DOI 없는 로컬 PDF가 같은 제목·연도로 여러 개면(같은 파일의 사본) 가장 작은 ID만 active로 둔다."""
    groups: Dict[Tuple[str, Any], List[Dict[str, Any]]] = {}
    for p in db["papers"]:
        if p.get("origin", "local") != "local" or p.get("status") in INACTIVE_STATUSES or normalize_doi(p.get("doi")):
            continue
        h = title_hash(p)
        if h:
            groups.setdefault((h, p.get("year")), []).append(p)
    n = 0
    for key, lst in groups.items():
        if len(lst) < 2:
            continue
        lst.sort(key=lambda x: as_int(x["id"]))
        keep = lst[0]
        for dup in lst[1:]:
            dup["status"] = "duplicate"
            dup["duplicate_of"] = keep["id"]
            dup["duplicate_reason"] = "same normalized title + year (no DOI)"
            dup["updated_at"] = now_iso()
            keep["detected_acupoints"] = sorted(set(as_list(keep.get("detected_acupoints")) + as_list(dup.get("detected_acupoints"))))
            n += 1
    return n


# --------------------------------------------------------------------------------------
# export-xlsx
# --------------------------------------------------------------------------------------
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def clean_text(v):
    """openpyxl이 거부하는 제어 문자 제거 (PDF 추출 텍스트에 흔함)."""
    if isinstance(v, str):
        return _ILLEGAL_XML.sub("", v)
    return v


def _ws_write(ws, header: List[str], rows: List[list], widths: Optional[List[int]] = None) -> None:
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([("" if v is None else clean_text(v)) for v in r])
    ws.freeze_panes = "A2"
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def meridian_sort_key(code: str):
    m = re.match(r"^([A-Z]+)(\d+)$", code)
    if not m:
        return (99, 0, code)
    pre = m.group(1)
    return (MERIDIAN_ORDER.index(pre) if pre in MERIDIAN_ORDER else 98, int(m.group(2)), code)


def cmd_export_xlsx(ctx: Ctx, db: Dict[str, Any]) -> None:
    if openpyxl is None:
        print("[Error] openpyxl 미설치")
        return
    papers = [p for p in db["papers"] if p.get("status") != "superseded"]
    by_id = {as_int(p["id"]): p for p in papers}
    entries = acupoint_entries(db)
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = SHEET_MASTER

    rows1 = []
    for code in sorted(entries, key=meridian_sort_key):
        e = entries[code]
        if not (e.get("local_papers_count") or e.get("external_papers_count") or e.get("innervation_anatomy") or e.get("landmark_paper_ids")):
            continue
        lm_ids = as_list(e.get("landmark_paper_ids"))
        lm = by_id.get(lm_ids[0]) if lm_ids and lm_ids[0] in by_id else None
        lm_title = f"{lm.get('title','')} ({lm.get('journal','')}, {lm.get('year','')})" if lm else ""
        pmid_doi = f"PMID:{lm.get('pmid','')} / DOI:{lm.get('doi','')}" if lm else ""
        ev = (lm.get("evidence_summary") or lm.get("abstract_snippet") or "") if lm else ""
        ids = as_list(e.get("local_paper_ids"))
        rows1.append([
            code, e.get("name_kr", ""), e.get("name_cn", ""), e.get("meridian", ""), e.get("innervation_anatomy", ""),
            e.get("spinal_segment", ""), e.get("reflex_circuit", ""), lm_title, pmid_doi, ev,
            e.get("standard_anatomy_reference", ""), e.get("clinical_indications_mapped", ""), len(ids),
            (", ".join(map(str, ids[:8])) + (f" ... (총 {len(ids)}편)" if len(ids) > 8 else "")) if ids else "",
            e.get("external_papers_count", 0), e.get("last_api_sync_at", ""),
        ])
    _ws_write(ws1, ["경혈코드", "경혈명(한글)", "경혈명(한자)", "소속경락", "해부학적 신경·혈관 지배구조", "척수 분절", "생의학적 신경 반사 경로",
                    "대표 랜드마크 논문 (Title / Journal / Year)", "PMID / DOI", "실증 연구팀 및 핵심 기전 요약", "신경해부학 표준 근거",
                    "매핑된 임상 주치증", "소장논문수(편)", "소장 논문 ID 및 샘플", "외부수집논문수(편)", "최종API동기화"], rows1,
              [9, 9, 9, 22, 40, 14, 40, 40, 30, 40, 28, 30, 10, 36, 10, 20])

    ws2 = wb.create_sheet(SHEET_LOCAL_INDEX)
    rows2 = []
    for code in sorted(entries, key=meridian_sort_key):
        e = entries[code]
        ids = as_list(e.get("local_paper_ids"))
        if not ids:
            continue
        rows2.append([code, e.get("name_kr", ""), e.get("name_cn", ""), e.get("meridian", ""), len(ids),
                      ", ".join(map(str, ids)), by_id[ids[0]].get("filename", "") if ids[0] in by_id else ""])
    _ws_write(ws2, ["경혈코드", "경혈명(한글)", "경혈명(한자)", "소속경락", "소장 논문 수(편)", "매핑된 논문 ID 목록", "대표 소장 논문 파일명"], rows2,
              [9, 9, 9, 22, 10, 60, 50])

    ws3 = wb.create_sheet(SHEET_PAPERS)
    rows3 = []
    for p in sorted(papers, key=lambda x: as_int(x["id"])):
        rows3.append([as_int(p["id"]), p.get("filename") or p.get("url", ""), p.get("year"), p.get("journal", ""), p.get("authors", ""),
                      p.get("official_title") or p.get("title", ""), p.get("doi", ""), p.get("pmid", ""),
                      ", ".join(as_list(p.get("detected_acupoints"))), (p.get("abstract_snippet") or p.get("snippet") or "")[:250],
                      p.get("origin", "local"), p.get("status", "active"), p.get("content_hash", ""), (p.get("added_at") or "")[:10]])
    _ws_write(ws3, ["논문 ID", "파일명", "발행년도", "저널/출처", "저자", "논문 제목", "DOI", "PMID", "감지된 경혈 목록", "초록/본문 스니펫",
                    "출처(origin)", "상태", "콘텐츠해시", "추가일"], rows3, [8, 50, 8, 24, 20, 50, 28, 10, 20, 60, 10, 8, 14, 11])

    ws4 = wb.create_sheet(SHEET_EXTERNAL)
    rows4 = []
    for p in sorted([p for p in papers if p.get("origin", "local") != "local"], key=lambda x: as_int(x["id"])):
        rows4.append([as_int(p["id"]), p.get("source_db") or p.get("origin"), p.get("year"), p.get("journal", ""), p.get("title", ""),
                      p.get("authors", ""), p.get("doi", ""), p.get("pmid", ""), ", ".join(as_list(p.get("detected_acupoints"))),
                      p.get("language", ""), ", ".join(as_list(p.get("mesh")))[:200], (p.get("source_query") or "")[:200],
                      (p.get("fetched_at") or p.get("added_at") or "")[:19], p.get("url", ""), p.get("cn_title", "")])
    _ws_write(ws4, ["논문 ID", "수집 소스", "발행년도", "저널", "논문 제목", "저자", "DOI", "PMID", "매핑 경혈", "언어", "MeSH", "수집 쿼리",
                    "수집일시", "URL", "중국어 원제"], rows4, [8, 12, 8, 24, 50, 30, 28, 10, 14, 6, 30, 40, 19, 36, 30])

    ws5 = wb.create_sheet(SHEET_LOG)
    rows5 = []
    for ev in db["_metadata"].get("sync_log", []):
        rows5.append([ev.get("ts"), ev.get("trigger"), ev.get("action"), ev.get("source"), ev.get("acupoint", ""), (ev.get("query") or "")[:200],
                      ev.get("http_status"), ev.get("fetched"), ev.get("deduped"), ev.get("added"), ev.get("duration_ms"), ev.get("note", "")])
    _ws_write(ws5, ["일시", "트리거", "동작", "소스", "경혈", "쿼리", "HTTP상태", "수집건수", "중복제거", "신규적재", "소요(ms)", "비고"], rows5,
              [19, 8, 14, 10, 8, 50, 8, 8, 8, 8, 9, 40])

    if ctx.dry_run:
        print(f"[dry-run] xlsx 저장 생략 (시트1 {len(rows1)}행, 시트2 {len(rows2)}행, 시트3 {len(rows3)}행, 외부 {len(rows4)}행, 로그 {len(rows5)}행)")
        return
    backup_file(ctx, ctx.xlsx_path)
    tmp = ctx.xlsx_path.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    os.replace(tmp, ctx.xlsx_path)
    print(f"[export-xlsx] 저장: {ctx.xlsx_path} (시트1 {len(rows1)}행, 시트2 {len(rows2)}행, 시트3 {len(rows3)}행, 외부 {len(rows4)}행, 로그 {len(rows5)}행)")


# --------------------------------------------------------------------------------------
# validate
# --------------------------------------------------------------------------------------
def cmd_validate(ctx: Ctx, db: Dict[str, Any], check_xlsx: bool = True) -> List[str]:
    errs: List[str] = []
    papers = db["papers"]
    ids = [p.get("id") for p in papers]
    if any(not isinstance(i, int) for i in ids):
        errs.append("papers.id 중 정수가 아닌 값 존재")
    if len(set(ids)) != len(ids):
        errs.append("papers.id 중복")
    idset = set(ids)
    for c, lst in db["acupoint_inverted_index"].items():
        if std_code(c) != c:
            errs.append(f"역색인 코드 비표준: {c}")
        for i in lst:
            if i not in idset:
                errs.append(f"역색인 {c} → 존재하지 않는 ID {i}")
                break
    active_local = {p["file_path"] for p in papers if p.get("origin", "local") == "local" and p.get("status", "active") == "active"}
    man = set(db.get("file_manifest", {}).keys())
    if active_local - man:
        errs.append(f"manifest에 없는 active 로컬 레코드 {len(active_local - man)}건")
    dois, pmids = {}, {}
    for p in papers:
        if p.get("status") in INACTIVE_STATUSES:
            continue
        d = normalize_doi(p.get("doi"))
        if d:
            dois.setdefault(d, []).append(p["id"])
        pm = str(p.get("pmid") or "")
        if pm:
            pmids.setdefault(pm, []).append(p["id"])
    dd = {k: v for k, v in dois.items() if len(v) > 1}
    pd_ = {k: v for k, v in pmids.items() if len(v) > 1}
    if dd:
        errs.append(f"DOI 중복 {len(dd)}건 (예: {list(dd.items())[:2]})")
    if pd_:
        errs.append(f"PMID 중복 {len(pd_)}건 (예: {list(pd_.items())[:2]})")
    bad_types = sum(1 for p in papers if not isinstance(p.get("detected_acupoints"), list) or not isinstance(p.get("parsed_filename", {}), dict))
    if bad_types:
        errs.append(f"문자열화된 필드 보유 레코드 {bad_types}건")
    for code, e in acupoint_entries(db).items():
        if std_code(code) != code:
            errs.append(f"경혈 키 비표준: {code}")
            break
    if check_xlsx and ctx.xlsx_path.exists() and openpyxl:
        wb = openpyxl.load_workbook(ctx.xlsx_path, read_only=True)
        if SHEET_PAPERS in wb.sheetnames:
            n3 = sum(1 for _ in wb[SHEET_PAPERS].iter_rows(min_row=2, values_only=True))
            n_json = sum(1 for p in papers if p.get("status") != "superseded")
            if n3 != n_json:
                errs.append(f"시트 '{SHEET_PAPERS}' 행수 {n3} ≠ JSON {n_json}")
        else:
            errs.append(f"시트 '{SHEET_PAPERS}' 없음")
        if SHEET_LOCAL_INDEX in wb.sheetnames:
            for r in wb[SHEET_LOCAL_INDEX].iter_rows(min_row=2, values_only=True):
                code = str(r[0]).strip()
                e = db.get(code, {})
                if as_int(r[4]) != as_int(e.get("local_papers_count")):
                    errs.append(f"시트2 {code} 소장논문수 {r[4]} ≠ JSON {e.get('local_papers_count')}")
                    break
        else:
            errs.append(f"시트 '{SHEET_LOCAL_INDEX}' 없음")
        wb.close()
    if errs:
        print("[validate] FAIL")
        for e in errs:
            print("   -", e)
    else:
        print("[validate] OK")
    return errs


# --------------------------------------------------------------------------------------
# migrate-v3
# --------------------------------------------------------------------------------------
def cmd_migrate_v3(ctx: Ctx, db: Dict[str, Any]) -> None:
    md = db["_metadata"]
    if md.get("schema_version", 0) >= SCHEMA_VERSION:
        print("[migrate-v3] 이미 v3 스키마입니다. 건너뜀.")
        return
    print("[migrate-v3] 시작")
    base = local_dir(ctx, db)
    added_at = md.get("integrated_at") or now_iso()

    # 1) 경혈 키 정규화 + 문자열화 값 교정
    entries = acupoint_entries(db)
    merged_keys = 0
    for key in list(entries):
        e = db[key]
        for k in ("landmark_papers", "local_paper_ids", "local_papers_sample", "landmark_paper_ids", "external_paper_ids"):
            if k in e:
                e[k] = as_list(e[k])
        e["local_papers_count"] = as_int(e.get("local_papers_count"))
        e["code"] = std_code(e.get("code") or key)
        sc = std_code(key)
        if sc != key:
            if sc in db and isinstance(db[sc], dict) and "code" in db[sc]:
                tgt = db[sc]
                for k, v in e.items():
                    if (not tgt.get(k)) and v:
                        tgt[k] = v
                merged_keys += 1
            else:
                db[sc] = e
            del db[key]
    print(f"  경혈 키 정규화 완료 (병합 {merged_keys}건)")

    # 2) papers 타입 교정 · 로컬 필드 채움
    missing_files = 0
    for p in db["papers"]:
        p["id"] = as_int(p.get("id"))
        p["detected_acupoints"] = sorted({std_code(c) for c in as_list(p.get("detected_acupoints"))})
        p["parsed_filename"] = as_dict(p.get("parsed_filename"))
        p.setdefault("origin", "local")
        p.setdefault("status", "active")
        p["doi"] = normalize_doi(p.get("doi"))
        p["pmid"] = str(p.get("pmid") or "")
        meta = p["parsed_filename"]
        p.setdefault("title", meta.get("title") or re.sub(r"\.pdf$", "", p.get("filename", ""), flags=re.I))
        p.setdefault("authors", meta.get("author") or "")
        p.setdefault("journal", meta.get("journal") or "")
        p.setdefault("year", meta.get("year"))
        p.setdefault("volume_issue", str(meta.get("year") or ""))
        p["abstract_snippet"] = p.get("abstract_snippet") or p.get("snippet", "")
        p.setdefault("keywords", [])
        p.setdefault("mesh", [])
        p.setdefault("language", "")
        p.setdefault("evidence_level", None)
        p.setdefault("added_at", added_at)
        p["updated_at"] = now_iso()
        fp = p.get("file_path", "")
        abs_path = Path(fp) if os.path.isabs(fp) else base / fp
        p["file_path"] = relpath_norm(abs_path, base)
        p["ext"] = abs_path.suffix.lower()
        if os.path.exists(lp(abs_path)):
            st = os.stat(lp(abs_path))
            p["file_size"], p["mtime"], p["content_hash"] = st.st_size, round(st.st_mtime, 3), file_hash(abs_path)
            p.setdefault("text_extracted", bool(p.get("snippet")))
        else:
            p["status"] = "missing"
            p.setdefault("file_size", 0); p.setdefault("mtime", 0); p.setdefault("content_hash", "")
            missing_files += 1
    dups = mark_local_duplicates(db)
    print(f"  papers {len(db['papers'])}편 정규화 (해시 계산 완료, 유실 {missing_files}, 동일 DOI 중복 PDF {dups}편 duplicate 표시)")

    # 3) landmark_papers → 레지스트리 이관
    by_doi, by_pmid, by_hash = build_registry_index(db)
    by_id = {p["id"]: p for p in db["papers"]}
    nid = next_id(db)
    moved = merged = 0
    for code, e in acupoint_entries(db).items():
        lm_ids = as_list(e.get("landmark_paper_ids"))
        for lmp in as_list(e.get("landmark_papers")):
            if not isinstance(lmp, dict):
                continue
            src_db = str(lmp.get("source_db") or "")
            origin = ORIGIN_MAP.get(src_db.lower(), "manual")
            dup = find_duplicate(lmp, by_doi, by_pmid, by_hash)
            if dup is not None:
                merge_into(by_id[dup], {k: lmp.get(k) for k in ("pmid", "doi", "journal", "year", "volume_issue")}, source=origin)
                if code not in by_id[dup]["detected_acupoints"]:
                    by_id[dup]["detected_acupoints"] = sorted(set(by_id[dup]["detected_acupoints"] + [code]))
                if dup not in lm_ids:
                    lm_ids.append(dup)
                merged += 1
                continue
            if origin == "local":
                origin = "manual"  # 레지스트리에서 못 찾은 로컬 표기 → 수기 등록으로 간주
            rec = {
                "id": nid, "origin": origin, "status": "active", "title": lmp.get("title", ""), "authors": lmp.get("authors", ""),
                "journal": lmp.get("journal", ""), "year": as_int(lmp.get("year"), None), "volume_issue": lmp.get("volume_issue", ""),
                "doi": normalize_doi(lmp.get("doi")), "pmid": str(lmp.get("pmid") or ""), "abstract_snippet": "", "snippet": "",
                "detected_acupoints": [code], "keywords": [], "mesh": [], "language": "", "evidence_level": None,
                "source_db": src_db or "manual", "source_query": "", "fetched_at": added_at, "url": "", "cn_title": "", "cn_abstract": "",
                "also_found_in": [origin], "filename": "", "file_path": "", "parsed_filename": {},
                "evidence_summary": lmp.get("evidence_summary", ""), "added_at": added_at, "updated_at": now_iso(),
            }
            db["papers"].append(rec)
            by_id[nid] = rec
            if rec["doi"]:
                by_doi[rec["doi"]] = nid
            if rec["pmid"]:
                by_pmid[rec["pmid"]] = nid
            h = title_hash(rec)
            if h:
                by_hash[h] = nid
            lm_ids.append(nid)
            nid += 1
            moved += 1
        e["landmark_paper_ids"] = lm_ids
        e.setdefault("external_paper_ids", [])
        e.setdefault("external_papers_count", 0)
        e.setdefault("last_api_sync_at", "")
    print(f"  landmark_papers 이관: 신규 {moved}, 기존 병합 {merged}")

    # 4) manifest
    # manifest는 '디스크에 있는 파일' 단위로 추적한다(상태와 무관). duplicate 레코드도 포함해야 재등록되지 않는다.
    db["file_manifest"] = {p["file_path"]: {"id": p["id"], "size": p.get("file_size", 0), "mtime": p.get("mtime", 0), "hash": p.get("content_hash", "")}
                           for p in db["papers"] if p.get("origin") == "local" and p.get("status") in ("active", "duplicate")}

    # 5) 메타데이터
    md["version"] = DB_VERSION
    md["schema_version"] = SCHEMA_VERSION
    md["migrated_at"] = now_iso()
    md["source_local_dir"] = str(base)
    md.setdefault("last_local_scan", {})
    md.setdefault("last_api_sync", {})
    md["description"] = "국제 공인 신경반사 회로, 랜드마크 논문(PMID/DOI), 로컬 소장 논문 역색인 및 외부 API 수집 논문이 통합된 단일 마스터 데이터셋 (v3.0: papers 레지스트리 + 영구 ID + manifest 증분 동기화)"
    recompute_derived(db, ctx)
    log_event(ctx, db, "migrate-v3", source="local", http_status=None, fetched=len(db["papers"]), added=moved, deduped=merged,
              note=f"keys_merged={merged_keys} missing_files={missing_files}")
    print("[migrate-v3] 완료")


# --------------------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description="경혈 생의학 통합 마스터 DB 동기화 (v3.0)")
    ap.add_argument("--root", default=str(DEFAULT_ROOT), help="프로젝트 루트 (기본: 스크립트 위치에서 추론)")
    ap.add_argument("--dry-run", action="store_true", help="디스크에 쓰지 않음")
    ap.add_argument("--trigger", default="cli", choices=["cli", "rank6", "batch"])
    sub = ap.add_subparsers(dest="cmd", required=True)

    s1 = sub.add_parser("scan-local"); s1.add_argument("--full-rescan", action="store_true"); s1.add_argument("--enrich", action="store_true"); s1.add_argument("--enrich-limit", type=int, default=0, help="--enrich 시 상위 N편만 보강(0=전체)"); s1.add_argument("--skip-export", action="store_true")
    s2 = sub.add_parser("sync-api")
    for s in (s2,):
        s.add_argument("--code", required=True); s.add_argument("--name-kr", default=""); s.add_argument("--name-cn", default="")
        s.add_argument("--pinyin", default=""); s.add_argument("--meridian", default=""); s.add_argument("--max-papers", type=int, default=8)
        s.add_argument("--skip-export", action="store_true")
    sub.add_parser("export-xlsx")
    sub.add_parser("validate")
    s5 = sub.add_parser("all")
    s5.add_argument("--code", required=True); s5.add_argument("--name-kr", default=""); s5.add_argument("--name-cn", default="")
    s5.add_argument("--pinyin", default=""); s5.add_argument("--meridian", default=""); s5.add_argument("--max-papers", type=int, default=8)
    s5.add_argument("--enrich", action="store_true")
    sub.add_parser("migrate-v3")
    s6 = sub.add_parser("relink", help="기존 외부 논문의 경혈 연결을 초록 전문으로 재검사(지명·약어 오적중 제거) + 로컬 사본 중복 표시")
    s6.add_argument("--no-refetch", action="store_true", help="초록을 다시 받지 않고 저장된 제목·초록 앞부분만으로 검사")
    a = ap.parse_args()

    ctx = Ctx(Path(a.root), a.dry_run, a.trigger)
    if not ctx.json_path.exists():
        print(f"[Error] 마스터 JSON 없음: {ctx.json_path}")
        return 1
    db = load_db(ctx)
    print(f"[DB] {ctx.json_path.name} | version={db['_metadata'].get('version')} schema={db['_metadata'].get('schema_version', 2)} papers={len(db['papers'])} | dry_run={a.dry_run}")

    if a.cmd == "migrate-v3":
        cmd_migrate_v3(ctx, db)
        save_db(ctx, db)
        return 0

    if db["_metadata"].get("schema_version", 0) < SCHEMA_VERSION and a.cmd != "validate":
        print("[Error] v3 스키마가 아닙니다. 먼저 `migrate-v3`를 실행하세요.")
        return 1

    def finish(skip_export=False) -> int:
        save_db(ctx, db)
        if skip_export:
            return 0
        cmd_export_xlsx(ctx, db)
        errs = cmd_validate(ctx, db, check_xlsx=not ctx.dry_run)
        if errs and not ctx.dry_run:
            print("[rollback] 검증 실패 → 백업본 복원")
            for path in (ctx.json_path, ctx.xlsx_path):
                bk = sorted(ctx.backup_dir.glob(f"{path.stem}_*{path.suffix}"))
                if bk:
                    shutil.copy2(bk[-1], path)
            return 1
        return 0

    if a.cmd == "scan-local":
        cmd_scan_local(ctx, db, a.full_rescan, a.enrich, a.enrich_limit)
        return finish(a.skip_export)
    if a.cmd == "sync-api":
        res = cmd_sync_api(ctx, db, a.code, a.name_kr, a.name_cn, a.pinyin, a.meridian, a.max_papers)
        rc = finish(a.skip_export)
        return 2 if res.get("all_failed") else rc
    if a.cmd == "export-xlsx":
        cmd_export_xlsx(ctx, db)
        return 1 if cmd_validate(ctx, db, check_xlsx=not ctx.dry_run) else 0
    if a.cmd == "validate":
        return 1 if cmd_validate(ctx, db) else 0
    if a.cmd == "relink":
        n_dup = mark_local_title_duplicates(db)
        print(f"[relink] 로컬 사본(같은 제목·연도, DOI 없음) duplicate 표시 {n_dup}편")
        cmd_relink(ctx, db, refetch=not a.no_refetch)
        return finish()
    if a.cmd == "all":
        cmd_scan_local(ctx, db, False, a.enrich)
        res = cmd_sync_api(ctx, db, a.code, a.name_kr, a.name_cn, a.pinyin, a.meridian, a.max_papers)
        rc = finish()
        return 2 if res.get("all_failed") else rc
    return 0


if __name__ == "__main__":
    sys.exit(main())
