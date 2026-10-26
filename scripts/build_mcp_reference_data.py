# -*- coding: utf-8 -*-
"""
build_mcp_reference_data.py — MCP 서버가 표준 라이브러리만으로 읽을 수 있도록 참조 데이터를 JSON으로 내보낸다

만드는 파일
  data/mcp_reference_data.json : 경혈별 2대 교재 해부구조·문헌별 자침 깊이(푼)·기존 Tier 3 보고서 색인
  data/classics/index.json     : 4대 원전 정제 마크다운 목록(파일명·SHA-1) — 원격 서버가 GitHub API 없이 목록·변경을 확인

원본
  data/경혈학공통교재 경혈_해부구조_데이터셋_refine.xlsx, data/김재효 수업_경혈_해부구조_데이터셋.xlsx  (시트 '경혈별_해부구조_종합')
  data/침구문헌별 자침깊이 정리_AI 분석용.xlsx  ('0'=금침, 빈칸=미기술, 정수=푼)
  01_경혈별_해부생리_분석보고서/**/<코드>_*_보고서.md  (머리말의 ARI 표기 줄)
  data/classics/<원전>/*.md

사용 (저장소 루트에서):
  python scripts/build_mcp_reference_data.py          # 생성
  python scripts/build_mcp_reference_data.py --check  # 원본이 바뀌었는데 JSON이 옛것이면 종료코드 1
"""
import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
REF_JSON = DATA / "mcp_reference_data.json"
CLASSICS = DATA / "classics"
CLASSICS_INDEX = CLASSICS / "index.json"
TEXTBOOKS = {
    "공통교재": DATA / "경혈학공통교재 경혈_해부구조_데이터셋_refine.xlsx",
    "김재효": DATA / "김재효 수업_경혈_해부구조_데이터셋.xlsx",
}
DEPTH_XLSX = DATA / "침구문헌별 자침깊이 정리_AI 분석용.xlsx"
REPORT_DIR = ROOT / "01_경혈별_해부생리_분석보고서"


def norm_code(code):
    m = re.match(r"^([A-Z]{2})0*(\d+)$", str(code or "").strip().upper())
    return f"{m.group(1)}{int(m.group(2))}" if m else str(code or "").strip().upper()


def split_cell(v):
    return [s.strip() for s in str(v or "").splitlines() if s.strip()]


def load_textbooks():
    import openpyxl
    out = {}
    for label, path in TEXTBOOKS.items():
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["경혈별_해부구조_종합"]
        rows = ws.iter_rows(values_only=True)
        hdr = list(next(rows))
        col = {h: i for i, h in enumerate(hdr) if h}
        for r in rows:
            code = norm_code(r[col["경혈코드"]])
            if not code:
                continue
            entry = {}
            for cat in ("근육", "신경", "혈관"):
                ko = split_cell(r[col[f"{cat}_한글"]]) if f"{cat}_한글" in col else []
                en = split_cell(r[col[f"{cat}_영문"]]) if f"{cat}_영문" in col else []
                entry[cat] = [{"ko": k, "en": en[i] if i < len(en) else ""} for i, k in enumerate(ko)]
            out.setdefault(code, {})[label] = entry
    return out


def load_depths():
    import openpyxl
    ws = openpyxl.load_workbook(DEPTH_XLSX, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [h for h in rows[1]]
    skip = {"N0", "경혈부호", "경혈명", "갑을경 경혈 분류 및 배열순서"}
    out = {}
    for r in rows[2:]:
        if not r or not r[1]:
            continue
        code = norm_code(r[1])
        vals = {}
        for i, h in enumerate(hdr):
            if not h or h in skip:
                continue
            v = r[i] if i < len(r) else None
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            vals[str(h)] = v  # None = 미기술(blank)
        out[code] = vals
    return out


def load_reports():
    out = {}
    for md in sorted(REPORT_DIR.rglob("*_보고서.md")):
        if md.stem.endswith("_en"):
            continue
        codes = re.findall(r"[A-Z]{2}\d{2}", md.name.split("_")[0])
        head = md.read_text(encoding="utf-8", errors="replace").splitlines()[:12]
        title = next((l.lstrip("# ").strip() for l in head if l.startswith("#")), md.stem)
        ari = [re.sub(r"\s+", " ", l.lstrip("> ").strip())[:600] for l in head if l.startswith(">") and "ARI" in l]
        rel = md.relative_to(ROOT).as_posix()
        for c in codes:
            out.setdefault(norm_code(c), []).append({"path": rel, "title": title, "ari_lines": ari})
    return out


def classics_index():
    books = {}
    for d in sorted(p for p in CLASSICS.iterdir() if p.is_dir()):
        files = []
        for f in sorted(d.glob("*.md")):
            files.append({"name": f.name, "sha1": hashlib.sha1(f.read_bytes().replace(b"\r\n", b"\n")).hexdigest()})
        if files:
            books[d.name] = files
    return {"generated_by": "scripts/build_mcp_reference_data.py", "books": books}


def build_reference():
    tb, dp, rp = load_textbooks(), load_depths(), load_reports()
    codes = sorted(set(tb) | set(dp) | set(rp))
    return {
        "generated_by": "scripts/build_mcp_reference_data.py",
        "sources": {
            "textbooks": {k: v.name for k, v in TEXTBOOKS.items()},
            "depth": DEPTH_XLSX.name,
            "depth_rule": "값은 푼(分). 0=금침, null=미기술. 한대 1푼=2.3mm",
        },
        "acupoints": {c: {"textbooks": tb.get(c, {}), "depth_fen": dp.get(c, {}), "reports": rp.get(c, [])} for c in codes},
    }


def dump(obj):
    return json.dumps(obj, ensure_ascii=False, indent=1, sort_keys=True) + "\n"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="생성물이 원본과 맞는지 확인만 한다")
    args = ap.parse_args()
    targets = [(REF_JSON, dump(build_reference())), (CLASSICS_INDEX, dump(classics_index()))]
    stale = []
    for path, text in targets:
        cur = path.read_text(encoding="utf-8") if path.exists() else None
        if cur != text:
            stale.append(path)
            if not args.check:
                path.write_text(text, encoding="utf-8", newline="\n")
    for path in stale:
        print(("[STALE] " if args.check else "[WRITE] ") + path.relative_to(ROOT).as_posix())
    if not stale:
        print("[OK] 참조 데이터가 최신입니다.")
    sys.exit(1 if (args.check and stale) else 0)


if __name__ == "__main__":
    main()
